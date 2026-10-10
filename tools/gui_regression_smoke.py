from __future__ import annotations

"""Current-contract GUI regression smoke restored from the historical R1.0.10 suite."""

import csv
import gc
import json
import os
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import wave
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from PySide6.QtCore import QPoint, QThreadPool, Qt
from PySide6.QtGui import QKeySequence
from PySide6.QtTest import QTest
from PySide6.QtWidgets import (
    QBoxLayout,
    QApplication,
    QDialog,
    QFileDialog,
    QLabel,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QScrollArea,
    QToolButton,
)
from openpyxl import load_workbook

from streaming_manager.app_paths import AppPaths
from streaming_manager.ui_settings import open_ui_settings
from streaming_manager.constants import (
    APP_VERSION,
    AUCTION_SOUNDTRACK_LOOP_ONE_KEY,
    AUCTION_SOUNDTRACK_MEDIA_ID_KEY,
    AUCTION_SOUNDTRACK_MUTE_KEY,
    AUCTION_SOUNDTRACK_VOLUME_KEY,
    STATUS_NOT_PLAYED,
    STATUS_PLAYED,
    TIMER_OVERLAY_BACKGROUND_KEY,
    AUCTION_LOTS_OVERLAY_BACKGROUND_MEDIA_ID_KEY,
    WHEEL_SOUNDTRACK_MEDIA_ID_KEY,
    WHEEL_SOUNDTRACK_MUTE_KEY,
    WHEEL_SOUNDTRACK_VOLUME_KEY,
)
from streaming_manager.database import Database, DuplicateGameError, Game
from streaming_manager.media import (
    MEDIA_CATEGORY_SOUNDTRACK,
    MEDIA_CATEGORY_OVERLAY_BACKGROUNDS,
    MEDIA_STORAGE_MANAGED,
    MEDIA_STORAGE_EXTERNAL,
    media_asset_available,
)
from streaming_manager.position_policy import (
    POSITION_COLUMNS_SINGLE,
    POSITION_COLUMNS_START_CURRENT,
    position_columns_for_surface,
)
from streaming_manager.shared_xlsx import (
    SHARED_XLSX_HEADERS,
    main_games_rows,
    presentation_hash,
    position_projection_matches,
    read_shared_xlsx,
    state_hash,
    write_shared_xlsx,
)
from streaming_manager.public_xlsx import (
    PUBLIC_XLSX_HEADERS,
    public_mirror_rows,
    public_state_hash,
    write_public_xlsx,
)
from streaming_manager.ui import MainWindow
from streaming_manager.views.games import DeleteAllGamesDialog, GameDialog
from tools.position_policy_smoke import assert_elimination_lifecycle



def assert_p08_background_media(app, window, root: Path) -> None:
    """Both selectors share D26 dedup/availability with no new registry."""
    db = window.db
    stream = window.stream_tab
    category = MEDIA_CATEGORY_OVERLAY_BACKGROUNDS
    source = root / "P08-Main.png"
    source.write_bytes(b"P08 test PNG fixture")
    with patch.object(QMessageBox, "information") as notice:
        stream._apply_background_import(source, MEDIA_STORAGE_MANAGED, "main")
        matches = db.media_assets_by_filename(category, source.name)
        if len(matches) != 1 or matches[0].storage_mode != MEDIA_STORAGE_MANAGED:
            raise AssertionError("P08 initial managed background was not registered")
        managed_id = matches[0].id
        stream._apply_background_import(source, MEDIA_STORAGE_EXTERNAL, "main")
        matches = db.media_assets_by_filename(category, source.name)
        if len(matches) != 1 or matches[0].id != managed_id:
            raise AssertionError("P08 managed -> external switch created a duplicate")
        if stream.background_combo.currentData() != managed_id:
            raise AssertionError("P08 duplicate background was not reselected")
        if not notice.called:
            raise AssertionError("P08 duplicate background notice is missing")

        external_source = root / "P08-External.png"
        external_source.write_bytes(b"P08 external image fixture")
        stream._apply_background_import(external_source, MEDIA_STORAGE_EXTERNAL, "auction_lots")
        original = db.media_assets_by_filename(category, external_source.name)
        if len(original) != 1 or original[0].storage_mode != MEDIA_STORAGE_EXTERNAL:
            raise AssertionError("P08 external background import failed")
        preserved_id = original[0].id
        stream._apply_background_import(external_source, MEDIA_STORAGE_MANAGED, "auction_lots")
        promoted = db.media_assets_by_filename(category, external_source.name)
        if (
            len(promoted) != 1
            or promoted[0].id != preserved_id
            or promoted[0].storage_mode != MEDIA_STORAGE_MANAGED
            or not media_asset_available(db.path.parent, promoted[0])
        ):
            raise AssertionError("P08 external -> managed did not preserve the media ID")
        if stream.auction_lots_background_combo.currentData() != preserved_id:
            raise AssertionError("P08 Auction Lots selector lost promoted media ID")

        # Real video copy is asynchronous; the same row must be promoted after
        # the existing copy worker finishes (not replaced with a second row).
        video_source = root / "P08-Video.mp4"
        video_source.write_bytes(b"P08 video copy fixture")
        stream._apply_background_import(video_source, MEDIA_STORAGE_EXTERNAL, "main")
        video_id = db.media_assets_by_filename(category, video_source.name)[0].id
        stream._apply_background_import(video_source, MEDIA_STORAGE_MANAGED, "main")
        # Refreshing the catalog while the copy is in flight must not create
        # a second managed row before the existing external ID is promoted.
        stream._background_assets()
        pending_rows = db.media_assets_by_filename(category, video_source.name)
        if len(pending_rows) != 1 or pending_rows[0].id != video_id:
            raise AssertionError("P08 in-flight video copy registered a second row")
        deadline = time.monotonic() + 8
        while stream._background_copy_worker is not None:
            app.processEvents()
            if time.monotonic() > deadline:
                raise AssertionError("P08 video-copy worker did not finish")
            QTest.qWait(10)
        app.processEvents()
        video_rows = db.media_assets_by_filename(category, video_source.name)
        if (
            len(video_rows) != 1
            or video_rows[0].id != video_id
            or video_rows[0].storage_mode != MEDIA_STORAGE_MANAGED
        ):
            raise AssertionError("P08 asynchronous video promotion duplicated the row")

        # A disappearing managed file is pruned and any saved selection is reset.
        db.set_settings_bulk({"overlay_background_media_id": str(managed_id)})
        (stream.background_dir / source.name).unlink()
        stream.refresh()
        if db.get_media_asset(managed_id) is not None:
            raise AssertionError("P08 missing managed background row survived sync")
        if stream.background_combo.currentData():
            raise AssertionError("P08 missing managed background stayed selected")
        if db.get_setting("overlay_background_media_id", ""):
            raise AssertionError("P08 missing managed background setting was not reset")

        # A missing external reference survives in DB but is unavailable in
        # both selectors; its saved Auction Lots selection is cleared.
        missing_source = root / "P08-Missing.png"
        missing_source.write_bytes(b"P08 missing reference fixture")
        stream._apply_background_import(missing_source, MEDIA_STORAGE_EXTERNAL, "auction_lots")
        missing_id = db.media_assets_by_filename(category, missing_source.name)[0].id
        db.set_settings_bulk({
            AUCTION_LOTS_OVERLAY_BACKGROUND_MEDIA_ID_KEY: str(missing_id),
        })
        missing_source.unlink()
        stream.refresh()
        if db.get_media_asset(missing_id) is None:
            raise AssertionError("P08 missing external reference should remain registered")
        if stream.background_combo.findData(missing_id) >= 0:
            raise AssertionError("P08 unavailable external asset leaked into main choices")
        if stream.auction_lots_background_combo.findData(missing_id) >= 0:
            raise AssertionError("P08 unavailable external asset leaked into Auction Lots")
        if db.get_setting(AUCTION_LOTS_OVERLAY_BACKGROUND_MEDIA_ID_KEY, ""):
            raise AssertionError("P08 unavailable external lot selection stayed persisted")

        missing_source.write_bytes(b"P08 restored at the original source path")
        stream.refresh()
        if stream.background_combo.findData(missing_id) < 0:
            raise AssertionError("P08 restored external background did not reappear")
        stream._apply_background_import(missing_source, MEDIA_STORAGE_EXTERNAL, "main")
        restored = db.media_assets_by_filename(category, missing_source.name)
        if len(restored) != 1 or restored[0].id != missing_id:
            raise AssertionError("P08 re-added external background duplicated media ID")

        # Historical mixed-mode duplicates are not silently picked or grown.
        ambiguous_source = root / "P08-Ambiguous.png"
        ambiguous_source.write_bytes(b"P08 ambiguous source")
        managed_copy = stream.background_dir / ambiguous_source.name
        managed_copy.write_bytes(b"P08 direct managed copy")
        db.sync_managed_media_category(category)
        db.register_external_media_asset(category, ambiguous_source)
        before = db.media_assets_by_filename(category, ambiguous_source.name)
        selected_before = stream.background_combo.currentData()
        stream._apply_background_import(ambiguous_source, MEDIA_STORAGE_EXTERNAL, "main")
        after = db.media_assets_by_filename(category, ambiguous_source.name)
        if len(before) != 2 or len(after) != 2:
            raise AssertionError("P08 ambiguous duplicates were implicitly modified")
        if stream.background_combo.currentData() != selected_before:
            raise AssertionError("P08 ambiguous duplicate implicitly moved selection")

    if hasattr(stream, "repair_background_btn"):
        raise AssertionError("P08 obsolete background repair action remains visible")
    print("P08_MEDIA_DEDUP_AVAILABILITY_BOTH_SELECTORS=PASS")


def free_local_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


def make_wav(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with wave.open(str(path), "wb") as out:
        out.setnchannels(1)
        out.setsampwidth(2)
        out.setframerate(8000)
        out.writeframes(b"\x00\x00" * 800)


def cleanup_tree(path: Path) -> None:
    gc.collect()
    last_error = None
    for _ in range(10):
        try:
            shutil.rmtree(path)
            return
        except FileNotFoundError:
            return
        except OSError as exc:
            last_error = exc
            gc.collect()
            time.sleep(0.15)
    if last_error is not None:
        raise last_error





def assert_position_policy() -> None:
    max_amount = {"mode": "max_amount", "status": "running"}
    tie_wheel = {"mode": "max_amount", "status": "awaiting_wheel"}
    direct_standard = {"mode": "weighted_wheel", "wheel_format": "standard"}
    direct_elimination = {
        "mode": "weighted_wheel",
        "wheel_format": "elimination",
    }

    assert position_columns_for_surface("main_list", max_amount) == POSITION_COLUMNS_SINGLE
    assert position_columns_for_surface("public_list", max_amount) == POSITION_COLUMNS_SINGLE
    assert position_columns_for_surface("auction", None) == POSITION_COLUMNS_SINGLE
    assert position_columns_for_surface("auction_overlay", None) == POSITION_COLUMNS_SINGLE
    assert position_columns_for_surface("auction", max_amount) == POSITION_COLUMNS_START_CURRENT
    assert position_columns_for_surface("auction_overlay", tie_wheel) == POSITION_COLUMNS_START_CURRENT

    assert position_columns_for_surface("auction", direct_standard) == POSITION_COLUMNS_SINGLE
    assert position_columns_for_surface(
        "auction_overlay",
        direct_standard,
        [{"result": "winner"}],
    ) == POSITION_COLUMNS_SINGLE
    assert position_columns_for_surface("auction", direct_elimination) == POSITION_COLUMNS_SINGLE
    assert position_columns_for_surface(
        "auction",
        direct_elimination,
        [{"result": "elimination_selected"}],
    ) == POSITION_COLUMNS_START_CURRENT
    assert position_columns_for_surface(
        "auction_overlay",
        {"mode": "weighted_wheel", "wheel_format": "standard"},
        [{"result": "eliminated"}],
    ) == POSITION_COLUMNS_START_CURRENT


def assert_shared_projection_preserves_external_edit(app, window, root: Path) -> None:
    """A derived-only refresh must import an already delivered cloud edit first."""
    db = window.db
    auction = window.auction_tab
    path = root / "p07_shared_projection_race.xlsx"
    auction_id = db.create_auction_session("max_amount", 60)

    def wait_for_worker() -> None:
        deadline = time.monotonic() + 8
        while auction._shared_xlsx_worker is not None:
            app.processEvents()
            if time.monotonic() >= deadline:
                raise AssertionError("P07 Shared XLSX worker did not finish")
            QTest.qWait(10)
        app.processEvents()

    try:
        initial = write_shared_xlsx(db, path)
        auction._shared_xlsx_activate(
            path, persist=False, import_now=False,
            initial_hash=initial["hash"],
            initial_presentation_hash=initial["presentation_hash"],
            initial_signature=initial["signature"],
        )
        auction._shared_xlsx_poll_timer.stop()
        workbook = load_workbook(path)
        edited_title = str(workbook.active["B2"].value)
        cloud_review = "P07 cloud review delivered before position refresh"
        workbook.active["G2"] = cloud_review
        workbook.save(path)
        workbook.close()
        top_points = max(game.sm_points for game in db.auction_eligible_games())
        db.add_or_increment_auction_lot(
            auction_id, "P07 Shared Temporary", int(top_points) + 10_000,
        )
        local_rows = main_games_rows(db)
        if state_hash(local_rows) != initial["hash"]:
            raise AssertionError("P07 race fixture changed logical permanent data")
        if presentation_hash(local_rows) == initial["presentation_hash"]:
            raise AssertionError("P07 race fixture did not change derived positions")

        auction.shared_xlsx_local_data_changed()
        auction._shared_xlsx_write_timer.stop()
        auction._shared_xlsx_start_local_write()
        wait_for_worker()
        preserved = read_shared_xlsx(path)
        if not any(
            row["title"] == edited_title and row["review"] == cloud_review
            for row in preserved["rows"]
        ):
            raise AssertionError("P07 position-only refresh overwrote a cloud business edit")

        # Reuse the existing stable-file poll/import path, then drain its
        # projection rewrite without waiting for wall-clock poll intervals.
        auction._shared_xlsx_poll()
        auction._shared_xlsx_poll()
        wait_for_worker()
        auction._shared_xlsx_write_timer.stop()
        auction._shared_xlsx_start_local_write()
        wait_for_worker()
        local_rows = main_games_rows(db)
        final_payload = read_shared_xlsx(path)
        if not any(
            row["title"] == edited_title and row["review"] == cloud_review
            for row in local_rows
        ):
            raise AssertionError("P07 cloud business edit was not imported into the database")
        if not position_projection_matches(local_rows, final_payload):
            raise AssertionError("P07 deferred Shared position projection did not converge")
        print("P07_SHARED_POSITION_REFRESH_PRESERVES_CLOUD_EDIT=PASS")
    finally:
        wait_for_worker()
        auction._shared_xlsx_disconnect()
        db.cancel_auction(auction_id)


def assert_position_xlsx_mirrors(db: Database, root: Path) -> None:
    shared_path = root / "p07_shared_position_mirror.xlsx"
    shared_result = write_shared_xlsx(db, shared_path)
    shared_rows = main_games_rows(db)
    shared_workbook = load_workbook(shared_path, read_only=True, data_only=True)
    try:
        sheet_rows = list(shared_workbook.active.iter_rows(values_only=True))
    finally:
        shared_workbook.close()
    if tuple(sheet_rows[0]) != SHARED_XLSX_HEADERS:
        raise AssertionError(f"Shared XLSX header mismatch: {sheet_rows[0]!r}")
    initial_import = read_shared_xlsx(shared_path)
    if not position_projection_matches(shared_rows, initial_import):
        raise AssertionError("Shared XLSX derived positions do not match local projection")

    expected_shared = {
        str(row["title"]): row["position"]
        for row in shared_rows
    }
    actual_shared = {
        str(row[1]): row[0]
        for row in sheet_rows[1:]
    }
    if actual_shared != expected_shared:
        raise AssertionError(
            f"Shared XLSX position mismatch: {actual_shared!r} != {expected_shared!r}"
        )
    if "UI076 Archive" not in actual_shared or "UI076 Temporary" in actual_shared:
        raise AssertionError("Shared XLSX position mirror included the wrong fixture scope")

    position_only_rows = [dict(row) for row in shared_rows]
    position_only_rows[0]["position"] = (
        int(position_only_rows[0]["position"] or 0) + 50
    )
    if state_hash(position_only_rows) != shared_result["hash"]:
        raise AssertionError("Shared logical hash included derived position")
    if presentation_hash(position_only_rows) == shared_result["presentation_hash"]:
        raise AssertionError("Shared presentation hash ignored derived position")

    edited = load_workbook(shared_path)
    edited.active["A2"] = 987654
    edited.save(shared_path)
    edited.close()
    imported = read_shared_xlsx(shared_path)
    if imported["hash"] != shared_result["hash"]:
        raise AssertionError("Manual Shared XLSX position edit changed import hash")
    if position_projection_matches(shared_rows, imported):
        raise AssertionError("Manual Shared XLSX position edit was not detected for rewrite")

    legacy_path = root / "p07_shared_legacy_headers.xlsx"
    shutil.copyfile(shared_path, legacy_path)
    legacy = load_workbook(legacy_path)
    legacy.active.delete_cols(1)
    legacy.save(legacy_path)
    legacy.close()
    legacy_import = read_shared_xlsx(legacy_path)
    if legacy_import["hash"] != shared_result["hash"]:
        raise AssertionError("Legacy Shared XLSX without position no longer imports")
    if legacy_import["has_position_column"]:
        raise AssertionError("Legacy Shared XLSX was mistaken for the new position projection")
    if position_projection_matches(shared_rows, legacy_import):
        raise AssertionError("Legacy Shared XLSX was not marked for position projection refresh")

    public_path = root / "p07_public_position_mirror.xlsx"
    public_result = write_public_xlsx(db, public_path)
    public_rows = public_mirror_rows(db)
    public_workbook = load_workbook(public_path, read_only=True, data_only=True)
    try:
        public_sheet_rows = list(public_workbook.active.iter_rows(values_only=True))
    finally:
        public_workbook.close()
    if tuple(public_sheet_rows[0]) != PUBLIC_XLSX_HEADERS:
        raise AssertionError(
            f"Public mirror XLSX header mismatch: {public_sheet_rows[0]!r}"
        )
    expected_public = {
        str(row["title"]): row["position"]
        for row in public_rows
    }
    actual_public = {
        str(row[1]): row[0]
        for row in public_sheet_rows[1:]
    }
    if actual_public != expected_public:
        raise AssertionError(
            f"Public XLSX position mismatch: {actual_public!r} != {expected_public!r}"
        )
    if {"UI076 Archive", "UI076 Temporary"} & set(actual_public):
        raise AssertionError("Public XLSX position mirror included an out-of-scope fixture")

    if public_rows:
        changed_public_rows = [dict(row) for row in public_rows]
        changed_public_rows[0]["position"] = (
            int(changed_public_rows[0]["position"] or 0) + 50
        )
        if public_state_hash(changed_public_rows) == public_result["hash"]:
            raise AssertionError("Public mirror hash ignored derived position")



def assert_wide_window_layout(app, window, test_sizes) -> None:
    # Wide-window manual QA: metadata filters stay adjacent after repeated
    # compact/wide transitions, for both populated and empty projections.
    games_tab = window.games_tab
    games_table = games_tab.table
    original_title_width = games_table.columnWidth(2)
    for filter_key in ("all", "archive"):
        games_tab.apply_stat_filter(filter_key)
        expected_rows = tuple(
            games_table.item(row, 2).text()
            for row in range(games_table.rowCount())
        )
        for width, height in test_sizes:
            window.resize(width, height)
            app.processEvents()
            QTest.qWait(150)
            app.processEvents()
            if window.width() != width or window.height() != height:
                raise AssertionError(
                    f"wide-window resize refused {width}x{height}: {window.size()}"
                )
            actual_rows = tuple(
                games_table.item(row, 2).text()
                for row in range(games_table.rowCount())
            )
            if actual_rows != expected_rows:
                raise AssertionError("resizing changed the main-list projection")
            if width == 2560 and games_tab._stats_columns != len(games_tab._stat_filter_buttons_order):
                raise AssertionError("filters did not return to one row at 2560px")
            if games_tab._stats_columns == len(games_tab._stat_filter_buttons_order):
                filters = games_tab._stat_filter_buttons_order
                for previous, following in zip(filters, filters[1:]):
                    previous_right = previous.mapTo(
                        games_tab, QPoint(previous.width(), 0)
                    ).x()
                    following_left = following.mapTo(games_tab, QPoint(0, 0)).x()
                    gap = following_left - previous_right
                    if not 4 <= gap <= 16:
                        raise AssertionError(
                            f"wide-window filters spread at {width}px: gap={gap}"
                        )
            else:
                for button in games_tab._stat_filter_buttons_order:
                    if button.width() < button.fontMetrics().horizontalAdvance(button.text()) + 12:
                        raise AssertionError(
                            f"compact filter caption clipped after wide resize: "
                            f"platform={QApplication.platformName()}, size={window.size()}, "
                            f"text={button.text()!r}, actual={button.width()}, "
                            f"caption={button.fontMetrics().horizontalAdvance(button.text())}, "
                            f"hint={button.sizeHint().width()}, minimum={button.minimumSizeHint().width()}, "
                            f"viewport={window._main_tab_scroll_by_page[games_tab].viewport().width()}, "
                            f"grid={games_tab._stats_layout.geometry()}"
                        )

            # An attempted manual shrink must protect the title itself,
            # including an empty table, without widening the date column.
            date_width = games_table.columnWidth(3)
            required_title_width = games_tab._game_title_header_min_width()
            games_table.setColumnWidth(2, 1)
            if games_table.columnWidth(2) < required_title_width:
                raise AssertionError("main-list title header width floor was not enforced")
            if games_table.columnWidth(3) != date_width:
                raise AssertionError("title width floor resized the adjacent date column")
            games_table.setColumnWidth(2, original_title_width)
    games_tab.apply_stat_filter("all")


def main() -> int:
    assert_position_policy()
    app = QApplication.instance() or QApplication(sys.argv)
    app.setOrganizationName("Local Streaming Tools QA")
    app.setApplicationName(f"InOneLine 1.0.8 GUI Regression Core {os.getpid()}")

    root = Path(tempfile.mkdtemp(prefix="inone_gui_regression_"))
    paths = AppPaths.from_root(root)
    paths.ensure_runtime_dirs()
    settings = open_ui_settings(paths.data_dir, migrate_native=False)
    settings.clear()
    probe = QMainWindow()
    probe.setMinimumSize(1100, 700)
    probe.setGeometry(120, 120, 1100, 700)
    settings.setValue("main_window/geometry", probe.saveGeometry())
    settings.setValue("main_window/maximized", False)
    settings.setValue("main_window/tab_index", 0)
    obsolete_list_keys = (
        "lists/visible", "games/list_visible", "public/list_visible",
    )
    for key in obsolete_list_keys:
        settings.setValue(key, False)
    settings.sync()
    probe.deleteLater()

    window = None
    original_excepthook = sys.excepthook
    uncaught: list[tuple[type[BaseException], BaseException]] = []

    def capture(exc_type, exc_value, traceback):
        uncaught.append((exc_type, exc_value))
        original_excepthook(exc_type, exc_value, traceback)

    sys.excepthook = capture
    try:
        import_db = Database(root / "p03_import_aliases.db")
        neutral_csv = root / "p03_neutral_headers.csv"
        neutral_csv.write_text(
            "НАЗВАНИЕ;ДАТА;БАЛЛЫ\n"
            "P03 Neutral;27.08.2019;1500\n",
            encoding="utf-8",
        )
        neutral_result = import_db.import_csv(neutral_csv)
        if neutral_result["created"] != 1:
            raise AssertionError(f"neutral CSV import result mismatch: {neutral_result!r}")
        neutral_game = import_db.find_game_by_title("P03 Neutral")
        if neutral_game is None:
            raise AssertionError("neutral CSV title header was not accepted")
        if neutral_game.release_date != "2019-08-27" or neutral_game.amount != 1500:
            raise AssertionError(
                "neutral CSV field mapping mismatch: "
                f"{neutral_game.release_date!r}, {neutral_game.amount!r}"
            )

        legacy_csv = root / "p03_legacy_headers.csv"
        legacy_csv.write_text(
            "НАЗВАНИЕ ИГРЫ;ДАТА ВЫХОДА;БАЛЛЫ\n"
            "P03 Legacy;28.08.2019;1600\n",
            encoding="utf-8",
        )
        legacy_result = import_db.import_csv(legacy_csv)
        if legacy_result["created"] != 1:
            raise AssertionError(f"legacy CSV import result mismatch: {legacy_result!r}")
        legacy_game = import_db.find_game_by_title("P03 Legacy")
        if legacy_game is None:
            raise AssertionError("legacy CSV title alias stopped working")
        if legacy_game.release_date != "2019-08-28" or legacy_game.amount != 1600:
            raise AssertionError(
                "legacy CSV alias mapping mismatch: "
                f"{legacy_game.release_date!r}, {legacy_game.amount!r}"
            )

        legacy_pipe_csv = root / "p05_legacy_pipe.csv"
        legacy_pipe_csv.write_text("P05 Legacy Pipe|1700\n", encoding="utf-8")
        pipe_result = import_db.import_csv(legacy_pipe_csv)
        if pipe_result["created"] != 1:
            raise AssertionError(f"legacy pipe import result mismatch: {pipe_result!r}")
        pipe_game = import_db.find_game_by_title("P05 Legacy Pipe")
        if pipe_game is None or pipe_game.amount != 1700:
            raise AssertionError("legacy Название|Баллы import compatibility failed")

        def expect_csv_error(name, content, expected_parts, *, with_backup=False):
            csv_path = root / name
            csv_path.write_text(content, encoding="utf-8")
            try:
                if with_backup:
                    import_db.import_csv_with_backup(
                        csv_path,
                        root / "p03_error_backups",
                        merge=True,
                        keep_backups=3,
                    )
                else:
                    import_db.import_csv(csv_path)
            except Exception as exc:
                message = str(exc)
            else:
                raise AssertionError(f"CSV error case unexpectedly passed: {name}")
            for expected in expected_parts:
                if expected not in message:
                    raise AssertionError(
                        f"CSV error wording missing {expected!r} for {name}: {message!r}"
                    )
            return message

        expect_csv_error(
            "p03_duplicate.csv",
            "НАЗВАНИЕ;БАЛЛЫ\n"
            "P03 Duplicate;10\n"
            " p03 duplicate ;20\n",
            (
                "Повторяется название «p03 duplicate» (строка 3).",
                "Ранее это же название указано в строке 2 как «P03 Duplicate».",
                "Оставьте в CSV только одну запись с этим названием.",
            ),
        )

        expect_csv_error(
            "p03_invalid_status.csv",
            "НАЗВАНИЕ;СТАТУС\n"
            "P03 Bad Status;НЕИЗВЕСТНО\n",
            (
                "Некорректное значение в поле «СТАТУС»",
                "строка 2",
                "«НЕИЗВЕСТНО»",
                "ПРОХОДИТСЯ, НЕ ИГРАЛ, ИГРАЛ, ПРОЙДЕНО или ЗАБРОШЕНО",
            ),
        )

        expect_csv_error(
            "p03_invalid_coop.csv",
            "НАЗВАНИЕ;КООП/НЕ КООП\n"
            "P03 Bad Coop;ДА\n",
            (
                "Некорректное значение в поле «КООП/НЕ КООП»",
                "строка 2",
                "«ДА»",
                "Допустимо: КООП или НЕ КООП.",
            ),
        )

        expect_csv_error(
            "p03_invalid_date.csv",
            "НАЗВАНИЕ;ДАТА\n"
            "P03 Bad Date;31.02.2019\n",
            (
                "Некорректное значение в поле «ДАТА»",
                "строка 2",
                "«31.02.2019»",
                "например 27.08.2019",
                "или оставьте поле пустым.",
            ),
        )

        expect_csv_error(
            "p03_invalid_points.csv",
            "НАЗВАНИЕ;БАЛЛЫ\n"
            "P03 Bad Points;abc\n",
            (
                "Некорректное значение в поле «БАЛЛЫ»",
                "строка 2",
                "«abc»",
                "Укажите целое число; если баллов нет, укажите 0.",
            ),
        )

        missing_header_message = expect_csv_error(
            "p03_missing_title_header.csv",
            "БАЛЛЫ;ДАТА\n"
            "10;27.08.2019\n",
            (
                "Не найден обязательный столбец «НАЗВАНИЕ».",
                "старое «НАЗВАНИЕ ИГРЫ» также принимается",
                "«Название|Баллы»",
            ),
        )
        if "Название игры|Баллы" in missing_header_message:
            raise AssertionError(
                "missing-header error exposes legacy game-centric format wording"
            )

        expect_csv_error(
            "p03_legacy_missing_points.csv",
            "P03 Legacy Missing|\n",
            (
                "Не указаны баллы у «P03 Legacy Missing» (строка 1).",
                "значение после «|» обязательно",
                "если баллов нет, укажите 0.",
            ),
        )

        atomic_message = expect_csv_error(
            "p03_atomicity.csv",
            "НАЗВАНИЕ;СТАТУС\n"
            "P03 Atomic Good;НЕ ИГРАЛ\n"
            "P03 Atomic Bad;НЕИЗВЕСТНО\n",
            (
                "Некорректное значение в поле «СТАТУС»",
                "P03 Atomic Bad",
                "строка 3",
            ),
        )
        if import_db.find_game_by_title("P03 Atomic Good") is not None:
            raise AssertionError(
                f"CSV atomicity changed after error wording update: {atomic_message!r}"
            )

        backup_error = expect_csv_error(
            "p03_backup_path_error.csv",
            "НАЗВАНИЕ;БАЛЛЫ\n"
            "P03 Backup Error;abc\n",
            (
                "Импорт не выполнен:",
                "Некорректное значение в поле «БАЛЛЫ»",
                "Резервная копия перед попыткой импорта:",
            ),
            with_backup=True,
        )
        if "Резервная копия перед попыткой импорта:" not in backup_error:
            raise AssertionError("CSV error lost safety-backup path")

        db = Database(paths.database_path)
        db.set_setting("api_port", str(free_local_port()))

        first_id = db.add_game(
            Game(None, "Smoke Game A", "2026-01-01", 10_000, 0, STATUS_PLAYED, "")
        )
        db.add_game(
            Game(None, "Smoke Game B", "2026-02-01", 5_000, 0, STATUS_NOT_PLAYED, "")
        )

        wheel_track = root / "wheel-smoke.wav"
        auction_track = root / "auction-smoke.wav"
        make_wav(wheel_track)
        make_wav(auction_track)
        wheel_asset = db.register_external_media_asset(MEDIA_CATEGORY_SOUNDTRACK, wheel_track)
        auction_asset = db.register_external_media_asset(MEDIA_CATEGORY_SOUNDTRACK, auction_track)
        db.set_settings_bulk({
            WHEEL_SOUNDTRACK_MEDIA_ID_KEY: str(wheel_asset.id),
            WHEEL_SOUNDTRACK_VOLUME_KEY: "37",
            WHEEL_SOUNDTRACK_MUTE_KEY: "0",
            AUCTION_SOUNDTRACK_MEDIA_ID_KEY: str(auction_asset.id),
            AUCTION_SOUNDTRACK_VOLUME_KEY: "29",
            AUCTION_SOUNDTRACK_MUTE_KEY: "0",
            AUCTION_SOUNDTRACK_LOOP_ONE_KEY: "0",
            TIMER_OVERLAY_BACKGROUND_KEY: "color",
        })

        print("[1/4] Main window / cold-start layout")
        window = MainWindow(db, paths)
        window.show()
        app.processEvents()
        QTest.qWait(250)
        app.processEvents()

        if "--wide-window-probe" in sys.argv:
            assert_wide_window_layout(
                app, window,
                ((2560, 1440), (520, 640), (2560, 1440), (1100, 750)),
            )
            if uncaught:
                raise AssertionError(f"wide-window GUI exception: {uncaught[0]!r}")
            print("WIDE_WINDOW_2560x1440_GEOMETRY=PASS")
            return 0

        # P09 / UI-033: OBS help remains short, Russian and audio-aware.
        with patch.object(QMessageBox, "information") as show_obs_help:
            window.stream_tab._show_obs_widget_help()
        if show_obs_help.call_count != 1:
            raise AssertionError("P09 OBS help did not open its existing dialog")
        help_title, help_text = show_obs_help.call_args.args[1:3]
        if help_title != "Как добавить виджет в OBS":
            raise AssertionError("P09 OBS help title was changed")
        required_help = (
            "Задайте ширину и высоту",
            "музыкального плеера",
            "аукциона или колеса",
            "микшере OBS",
            "Отключать источник, когда он не виден",
            "Обновлять браузер при активации сцены",
        )
        if any(part not in help_text for part in required_help):
            raise AssertionError("P09 OBS help lost agreed Russian/audio guidance")
        if any(part in help_text for part in (
            "Width / Height", "Shutdown source when not visible",
            "Refresh browser when scene becomes active",
        )):
            raise AssertionError("P09 OBS help still exposes old English setting names")
        print("P09_UI033_OBS_HELP=PASS")

        # UI-034: the original title typography row belongs immediately
        # below Current Game, not in the later generic font section.
        stream = window.stream_tab
        form = stream.main_settings_form
        current_row, _ = form.getWidgetPosition(stream.game_combo)
        title_row, _ = form.getWidgetPosition(stream.title_typography_host)
        if title_row != current_row + 1 or form.rowCount() != 2:
            raise AssertionError(
                "P09 UI-034 title typography is not directly below Current Game"
            )
        title_label, title_controls_row = stream.typography_row_widgets["title"]
        if (
            title_label.parentWidget() is not stream.title_typography_host
            or title_controls_row.parentWidget() is not stream.title_typography_host
            or len(stream.typography_controls) != 6
        ):
            raise AssertionError("P09 UI-034 title widgets were lost or duplicated")
        font_combo, size_spin, color_btn = stream.typography_controls["title"]
        previous_title_size = size_spin.value()
        size_spin.setValue(previous_title_size + 1)
        with patch.object(QMessageBox, "information"):
            stream.save()
        if db.get_setting("overlay_font_title_size") != str(previous_title_size + 1):
            raise AssertionError("P09 UI-034 existing title font save key changed")
        stream.refresh()
        if size_spin.value() != previous_title_size + 1:
            raise AssertionError("P09 UI-034 title font did not reload")
        print("P09_UI034_TITLE_FONT_MOVED_WITH_EXISTING_SAVE=PASS")

        # UI-035: one information block owns the original five related rows.
        info_form = stream.info_form
        info_rows = [
            stream.info_field,
            stream.info_enabled,
            stream.info_position,
            stream.frame_color_row_widgets["info"],
            stream.info_typography_host,
        ]
        actual_info_rows = [
            info_form.getWidgetPosition(widget)[0] for widget in info_rows
        ]
        if actual_info_rows != list(range(5)):
            raise AssertionError(
                f"P09 UI-035 information controls are split: {actual_info_rows}"
            )
        position_label = info_form.labelForField(stream.info_position)
        frame_label = info_form.labelForField(
            stream.frame_color_row_widgets["info"]
        )
        if (
            stream.info_enabled.text() != "Показывать на оверлее"
            or position_label is None
            or position_label.text() != "Положение информационного блока:"
            or frame_label is None
            or frame_label.text() != "Рамка информационного блока:"
        ):
            raise AssertionError("P09 UI-035 accepted information labels changed")
        info_title, info_font_row = stream.typography_row_widgets["info"]
        if (
            info_title.parentWidget() is not stream.info_typography_host
            or info_font_row.parentWidget() is not stream.info_typography_host
            or info_title.text() != "Настройки шрифта информационного блока"
            or len(stream.typography_controls) != 6
        ):
            raise AssertionError("P09 UI-035 info font widgets were duplicated or lost")

        defaults = {
            "stream_info": "",
            "stream_info_enabled": "1",
            "overlay_info_position": "auto",
            "overlay_frame_info_color": "#FFFFFF",
            "overlay_font_info_size": "17",
            "overlay_font_info_color": "#FFFFFF",
        }
        originals = {
            key: db.get_setting(key, value)
            for key, value in defaults.items()
        }
        info_size = stream.typography_controls["info"][1]
        stream.info_field.setText("P09 UI-035 hidden text is kept")
        stream.info_position.setCurrentIndex(
            stream.info_position.findData("bottom_left")
        )
        stream._set_color_button(stream.frame_color_buttons["info"], "#12ABCD")
        stream._set_color_button(stream.typography_controls["info"][2], "#ABCDEF")
        info_size.setValue(23)
        stream.info_enabled.setChecked(False)
        app.processEvents()
        if (
            stream.info_field.isEnabled()
            or stream.info_position.isHidden() is False
            or stream.frame_color_row_widgets["info"].isHidden() is False
            or stream.info_typography_host.isHidden() is False
            or stream.info_enabled.isHidden()
        ):
            raise AssertionError("P09 UI-035 disabled info controls have wrong visibility")

        with patch.object(QMessageBox, "information"):
            stream.save()
        expected_info = {
            "stream_info": "P09 UI-035 hidden text is kept",
            "stream_info_enabled": "0",
            "overlay_info_position": "bottom_left",
            "overlay_frame_info_color": "#12ABCD",
            "overlay_font_info_size": "23",
            "overlay_font_info_color": "#ABCDEF",
        }
        for key, value in expected_info.items():
            if db.get_setting(key) != value:
                raise AssertionError(f"P09 UI-035 settings failed to save: {key}")
        stream.refresh()
        if (
            stream.info_enabled.isChecked()
            or stream.info_field.text() != expected_info["stream_info"]
            or stream.info_field.isEnabled()
            or str(stream.info_position.currentData()) != "bottom_left"
            or info_size.value() != 23
        ):
            raise AssertionError("P09 UI-035 existing controls did not reload")

        db.set_settings_bulk(originals)
        stream.refresh()
        print("P09_UI035_INFO_PANEL_GROUP_AND_EXISTING_SAVE=PASS")

        if APP_VERSION != "1.0.8":
            raise AssertionError(f"wrong app version: {APP_VERSION}")
        if window.minimumWidth() != 520 or window.minimumHeight() != 360:
            raise AssertionError(
                f"P06 technical minimum changed: {window.minimumWidth()}x{window.minimumHeight()}"
            )
        if window.width() < window.minimumWidth() or window.height() < window.minimumHeight():
            raise AssertionError(f"cold-start window below technical minimum: {window.size()}")

        expected_tabs = [
            "Список",
            "Публичный список",
            "Музыка",
            "Аукцион",
            "История аукционов",
            "Журнал",
            "Стрим / OBS",
            "Настройки",
        ]
        actual_tabs = [window.tabs.tabText(i) for i in range(window.tabs.count())]
        if actual_tabs != expected_tabs:
            raise AssertionError(f"main tabs mismatch: {actual_tabs}")

        # UI-076: one existing aggregate serves both the statistics snapshot
        # and the right-side informational QLabel, never the visible rows.
        points_label = window.games_tab.total_points_label
        if not isinstance(points_label, QLabel):
            raise AssertionError("UI-076 total points must be an informational QLabel")
        if points_label in window.games_tab.filter_buttons.values():
            raise AssertionError("UI-076 total points was registered as a filter")
        if points_label.text() != "Всего баллов: 15000":
            raise AssertionError(f"UI-076 initial sum wrong: {points_label.text()!r}")
        if db.game_stats()["total_points"] != 15000:
            raise AssertionError("UI-076 base DB aggregate wrong")
        if hasattr(window.public_tab, "total_points_label"):
            raise AssertionError("UI-076 points counter leaked into Public List")

        # UI-076 stable position 2: the informational label follows the
        # existing sorting/URL/preview buttons. Repeated resizes must never
        # shift it horizontally (the former geometry-based spacer jittered).
        initial_ui076_size = window.size()
        from PySide6.QtCore import QPoint

        measured_action_gap = None
        measured_badge_left = None
        available_width = window.screen().availableGeometry().width()
        layout_max_width = min(1600, max(600, available_width - 32))
        if layout_max_width >= 1100:
            layout_test_widths = (1100, 1150, 1300, 1600, 1450, 1200, 1100)
        else:
            layout_test_widths = tuple(
                max(600, int(layout_max_width * fraction))
                for fraction in (0.72, 0.80, 0.90, 1.0, 0.85, 0.75, 1.0)
            )
        for width in layout_test_widths:
            window.resize(width, 750)
            app.processEvents()
            QTest.qWait(30)
            app.processEvents()
            if window.width() != width:
                raise AssertionError(
                    f"UI-076 resize refused {width}px: actual={window.width()}"
                )
            for button in window.games_tab.filter_buttons.values():
                caption_width = button.fontMetrics().horizontalAdvance(button.text())
                if button.width() < caption_width + 12:
                    raise AssertionError(
                        f"UI-076 filter clipped at {width}px: "
                        f"{button.text()!r} (width={button.width()}, text={caption_width})"
                    )
            if points_label.width() < points_label.sizeHint().width() - 1:
                raise AssertionError(
                    f"UI-076 points label clipped at {width}px: "
                    f"{points_label.width()} vs {points_label.sizeHint().width()}"
                )
            if window.games_tab._compact_controls_enabled:
                # At narrow widths the UI-076 badge joins the right-hand
                # vertical action stack, after preview and before clear-list.
                compact_stack = (
                    window.games_tab.sorting_rules_btn,
                    window.games_tab.copy_list_overlay_url_btn,
                    window.games_tab.open_list_overlay_preview_btn,
                    points_label,
                    window.games_tab.clear_all_btn,
                )
                stack_positions = [
                    (
                        widget.mapTo(window.games_tab, QPoint(0, 0)).y(),
                        widget.mapTo(
                            window.games_tab, widget.rect().topRight()
                        ).x(),
                    )
                    for widget in compact_stack
                ]
                if any(
                    stack_positions[index + 1][0] <= stack_positions[index][0]
                    for index in range(len(stack_positions) - 1)
                ):
                    raise AssertionError(
                        f"P06 compact actions not stacked in order at {width}px"
                    )
                stack_right_edges = [position[1] for position in stack_positions]
                if max(stack_right_edges) - min(stack_right_edges) > 3:
                    raise AssertionError(
                        f"P06 compact actions are not right-aligned at {width}px: "
                        f"{stack_right_edges}"
                    )
            else:
                sorting_y = window.games_tab.sorting_rules_btn.mapTo(
                    window.games_tab, QPoint(0, 0)
                ).y()
                badge_y = points_label.mapTo(
                    window.games_tab, QPoint(0, 0)
                ).y()
                if abs(sorting_y - badge_y) > 3:
                    raise AssertionError(
                        f"UI-076 points not on sorting actions row at {width}px: "
                        f"{sorting_y} vs {badge_y}"
                    )
                last_action_right = window.games_tab.open_list_overlay_preview_btn.mapTo(
                    window.games_tab,
                    window.games_tab.open_list_overlay_preview_btn.rect().topRight()
                ).x()
                badge_left = points_label.mapTo(
                    window.games_tab, QPoint(0, 0)
                ).x()
                gap = badge_left - last_action_right - 1
                if not 4 <= gap <= 16:
                    raise AssertionError(
                        f"UI-076 points not immediately after preview at {width}px: "
                        f"gap={gap}"
                    )
                if measured_action_gap is None:
                    measured_action_gap = gap
                    measured_badge_left = badge_left
                if abs(gap - measured_action_gap) > 1:
                    raise AssertionError(
                        f"UI-076 badge jitters relative to preview at {width}px: "
                        f"{gap} vs {measured_action_gap}"
                    )
                if abs(badge_left - measured_badge_left) > 1:
                    raise AssertionError(
                        f"UI-076 badge position shifts with window width at {width}px: "
                        f"{badge_left} vs {measured_badge_left}"
                    )
            if points_label.text() != "Всего баллов: 15000":
                raise AssertionError("UI-076 resizing changed the total")

        # Native Windows sizes stay within the runner's actual desktop.
        # An isolated offscreen process also exercises the exact 2560x1440 size.
        native_screen = window.screen().availableGeometry()
        native_width = max(520, min(2560, native_screen.width() - 80))
        native_height = max(360, min(1440, native_screen.height() - 80))
        assert_wide_window_layout(
            app, window,
            ((native_width, native_height), (520, min(640, native_height)),
             (native_width, native_height)),
        )
        subprocess.run(
            [sys.executable, str(Path(__file__).resolve()), "--wide-window-probe"],
            env={**os.environ, "QT_QPA_PLATFORM": "offscreen"},
            check=True,
            timeout=60,
        )
        window.resize(initial_ui076_size)
        app.processEvents()

        archived_points_id = db.add_game(
            Game(None, "UI076 Archive", "2026-02-11", 321, 0, STATUS_NOT_PLAYED, "")
        )
        db.archive_game(archived_points_id, True)
        temporary_points_id = db.add_game(
            Game(None, "UI076 Temporary", "2026-02-12", 7777, 0, STATUS_PLAYED, "")
        )
        # Fixture-only synthetic unmaterialized lot. Normal product data must
        # exclude auction_only=1 even when that row carries SM points.
        with db.connect() as conn:
            conn.execute(
                "UPDATE games SET auction_only=1 WHERE id=?", (temporary_points_id,)
            )
        window.games_tab.refresh()
        app.processEvents()
        if points_label.text() != "Всего баллов: 15321":
            raise AssertionError("UI-076 archive or temporary exclusion wrong")

        assert_position_xlsx_mirrors(db, root)

        public_export_expected = db.public_games()
        public_export_titles = {str(row["title"]) for row in public_export_expected}
        if {"UI076 Archive", "UI076 Temporary"} & public_export_titles:
            raise AssertionError("public_games included archived or auction_only fixtures")
        public_export_headers = ["НАЗВАНИЕ ИГРЫ", "БАЛЛЫ", "ОТЗЫВ", "СТАТУС"]

        def click_public_export(button, basename, suffix):
            selected_base = root / basename
            with (
                patch.object(
                    QFileDialog,
                    "getSaveFileName",
                    return_value=(str(selected_base), ""),
                ),
                patch.object(QMessageBox, "information") as info_box,
                patch.object(QMessageBox, "critical") as error_box,
            ):
                button.click()
            if error_box.called:
                raise AssertionError(
                    f"Public export {suffix} raised an error: {error_box.call_args!r}"
                )
            if not info_box.called:
                raise AssertionError(f"Public export {suffix} did not report success")
            saved_path = selected_base.with_suffix(suffix)
            if not saved_path.is_file():
                raise AssertionError(
                    f"Public export did not normalize its {suffix} suffix: {saved_path}"
                )
            return saved_path

        public_csv_path = click_public_export(
            window.public_tab.public_export_csv_btn, "p05_public_export_csv", ".csv"
        )
        with public_csv_path.open(encoding="utf-8-sig", newline="") as exported_file:
            public_csv_rows = list(csv.reader(exported_file, delimiter=";"))
        expected_public_csv_rows = [
            [
                str(row["title"]),
                str(int(row["sm_points"])),
                row["review"] or "",
                row["status"],
            ]
            for row in public_export_expected
        ]
        if public_csv_rows != [public_export_headers, *expected_public_csv_rows]:
            raise AssertionError(f"Public CSV contract mismatch: {public_csv_rows!r}")

        public_json_path = click_public_export(
            window.public_tab.public_export_json_btn, "p05_public_export_json", ".json"
        )
        public_json_payload = json.loads(public_json_path.read_text(encoding="utf-8"))
        if public_json_payload["columns"] != public_export_headers:
            raise AssertionError(
                f"Public JSON columns mismatch: {public_json_payload['columns']!r}"
            )
        if public_json_payload["games"] != public_export_expected:
            raise AssertionError("Public JSON changed the canonical public dataset")
        if any(
            key in {"СТАРТ", "ТЕКУЩАЯ"}
            for row in public_json_payload["games"]
            for key in row
        ):
            raise AssertionError("Public JSON added screen-only position columns")

        public_xlsx_path = click_public_export(
            window.public_tab.public_export_xlsx_btn, "p05_public_export_xlsx", ".xlsx"
        )
        public_xlsx_workbook = load_workbook(
            public_xlsx_path, read_only=True, data_only=True
        )
        try:
            public_xlsx_rows = list(
                public_xlsx_workbook.active.iter_rows(values_only=True)
            )
        finally:
            public_xlsx_workbook.close()
        expected_public_xlsx_rows = [
            tuple(public_export_headers),
            *[
                (row["title"], int(row["sm_points"]), row["review"] or None, row["status"])
                for row in public_export_expected
            ],
        ]
        if public_xlsx_rows != expected_public_xlsx_rows:
            raise AssertionError(f"Public XLSX contract mismatch: {public_xlsx_rows!r}")

        window.games_tab.apply_stat_filter(STATUS_PLAYED)
        window.games_tab.search.setText("UI076 Archive")
        app.processEvents()
        if points_label.text() != "Всего баллов: 15321":
            raise AssertionError("UI-076 total changed with active search/filter")
        if window.games_tab.active_filter != STATUS_PLAYED:
            raise AssertionError("UI-076 search unexpectedly changed statistic filter")
        window.games_tab.search.clear()
        db.update_game(archived_points_id, {"sm_points": 654})
        window.games_tab.refresh()
        app.processEvents()
        if points_label.text() != "Всего баллов: 15654":
            raise AssertionError("UI-076 archived points edit not reflected")

        # Simulate materialization: the formerly temporary row becomes a
        # normal record and must immediately participate in the same aggregate.
        with db.connect() as conn:
            conn.execute(
                "UPDATE games SET auction_only=0 WHERE id=?", (temporary_points_id,)
            )
        window.games_tab.refresh()
        app.processEvents()
        if points_label.text() != "Всего баллов: 23431":
            raise AssertionError("UI-076 materialized points missing")

        db.delete_game(temporary_points_id)
        db.delete_game(archived_points_id)
        window.games_tab.apply_stat_filter("all")
        window.games_tab.refresh()
        app.processEvents()
        if points_label.text() != "Всего баллов: 15000":
            raise AssertionError("UI-076 delete/cleanup failed to restore points")

        public_descriptions = [
            label.text()
            for label in window.public_tab.findChildren(QLabel)
            if label.text().startswith("Публичный список формируется напрямую")
        ]
        if len(public_descriptions) != 1:
            raise AssertionError(
                f"public-list description lookup mismatch: {public_descriptions!r}"
            )
        public_description = public_descriptions[0]
        if "ПОЗИЦИЯ / НАЗВАНИЕ / БАЛЛЫ / ОТЗЫВ / СТАТУС" not in public_description:
            raise AssertionError(
                f"public-list description title label mismatch: {public_description!r}"
            )
        if "НАЗВАНИЕ ИГРЫ" in public_description:
            raise AssertionError("legacy public-list description title label remains")
        public_position_header = window.public_tab.table.horizontalHeaderItem(0)
        if public_position_header is None or public_position_header.text() != "ПОЗИЦИЯ":
            raise AssertionError("public-list position header mismatch")
        public_title_header = window.public_tab.table.horizontalHeaderItem(1)
        if public_title_header is None or public_title_header.text() != "НАЗВАНИЕ":
            raise AssertionError(
                "public-list table title header mismatch: "
                f"{public_title_header.text() if public_title_header else None!r}"
            )

        public_button_texts = {
            button.text() for button in window.public_tab.findChildren(QPushButton)
        }
        if "Открыть локальный JSON" in public_button_texts:
            raise AssertionError("legacy Public List local JSON button remains")

        if window.games_tab.add_btn.text() != "Добавить":
            raise AssertionError(
                f"add button label mismatch: {window.games_tab.add_btn.text()!r}"
            )

        games_position_header = window.games_tab.table.horizontalHeaderItem(1)
        if games_position_header is None or games_position_header.text() != "ПОЗИЦИЯ":
            raise AssertionError("main-list position header mismatch")
        games_title_header = window.games_tab.table.horizontalHeaderItem(2)
        if games_title_header is None or games_title_header.text() != "НАЗВАНИЕ":
            raise AssertionError(
                "main-list title header mismatch: "
                f"{games_title_header.text() if games_title_header else None!r}"
            )
        # UI-008: the redundant main-list Find button is removed, while
        # Public's separate UI-028 decision is intentionally untouched.
        if hasattr(window.games_tab, "search_btn") or "Найти" in {
            button.text() for button in window.games_tab.findChildren(QPushButton)
        }:
            raise AssertionError("UI-008 main-list Find button still exists")
        # UI-009 removes only the main-list reset control, never the
        # native QLineEdit clear icon or statistic-filter toggle buttons.
        if hasattr(window.games_tab, "reset_filters_btn") or "Сбросить фильтры" in {
            button.text() for button in window.games_tab.findChildren(QPushButton)
        }:
            raise AssertionError("UI-009 obsolete Reset Filters button still exists")
        if not window.games_tab.search.isClearButtonEnabled():
            raise AssertionError("UI-009 native QLineEdit clear button disabled")
        # UI-028: Public no longer duplicates native live-search/Enter with
        # a visible Find action; the main list was retired by UI-008.
        if hasattr(window.public_tab, "search_btn") or "Найти" in {
            button.text() for button in window.public_tab.findChildren(QPushButton)
        }:
            raise AssertionError("UI-028 Public Find button still exists")
        if not window.public_tab.search.isClearButtonEnabled():
            raise AssertionError("UI-028 native Public search clear disabled")

        # P06: the approved order and small technical minimum keep the shell usable.
        expected_tab_titles = [
            "Список",
            "Публичный список",
            "Музыка",
            "Аукцион",
            "История аукционов",
            "Журнал",
            "Стрим / OBS",
            "Настройки",
        ]
        actual_tab_titles = [
            window.tabs.tabText(index) for index in range(window.tabs.count())
        ]
        if actual_tab_titles != expected_tab_titles:
            raise AssertionError(f"P06 main tab order mismatch: {actual_tab_titles!r}")
        if window.minimumWidth() >= 800 or window.minimumHeight() >= 600:
            raise AssertionError("P06 main-window minimum is still tied to content size")
        main_tab_pages = (
            window.games_tab,
            window.public_tab,
            window.music_tab,
            window.auction_tab,
            window.completed_history_tab,
            window.log_tab,
            window.stream_tab,
            window.settings_tab,
        )
        for page in main_tab_pages:
            scroll = window._main_tab_scroll_by_page[page]
            if not isinstance(scroll, QScrollArea):
                raise AssertionError("P06 workspace is missing its scroll host")
            if scroll.horizontalScrollBarPolicy() != Qt.ScrollBarAlwaysOff:
                raise AssertionError("P06 workspace-level horizontal scrolling is enabled")
            if scroll.verticalScrollBarPolicy() != Qt.ScrollBarAsNeeded:
                raise AssertionError("P06 workspace vertical scrolling is unavailable")
        if window.tabs.usesScrollButtons():
            raise AssertionError("P06 tab strip still displays horizontal navigation arrows")

        # Old installations persisted numeric indices in the former tab order.
        legacy_pages = [
            ("list", window.games_tab),
            ("public", window.public_tab),
            ("stream", window.stream_tab),
            ("music", window.music_tab),
            ("auction", window.auction_tab),
            ("auction_history", window.completed_history_tab),
            ("log", window.log_tab),
            ("settings", window.settings_tab),
        ]
        for legacy_index, (expected_key, expected_page) in enumerate(legacy_pages):
            window._ui_settings.remove("main_window/tab_key")
            window._ui_settings.setValue("main_window/tab_index", legacy_index)
            window._ui_settings.sync()
            window._restore_ui_state()
            app.processEvents()
            if window._current_main_page() is not expected_page:
                raise AssertionError(
                    f"P06 legacy tab index {legacy_index} restored the wrong page"
                )
            if window._ui_settings.value("main_window/tab_key", "", type=str) != expected_key:
                raise AssertionError(
                    f"P06 legacy tab index {legacy_index} did not save its stable key"
                )
            if window._ui_settings.value("main_window/tab_index", -1, type=int) != window.tabs.indexOf(window._main_tab_scroll_by_page[expected_page]):
                raise AssertionError(
                    f"P06 legacy tab index {legacy_index} did not migrate to the new index"
                )

        # P04-A: upgrading from hidden-list preferences must still show both
        # tables, preserve window state and keep explicit Enter navigation.
        original_size = window.size()
        if window.tabs.currentIndex() != window.tabs.indexOf(window._main_tab_scroll_by_page[window.settings_tab]):
            raise AssertionError("P06 legacy-index fixture did not finish at Settings")
        for tab in (window.games_tab, window.public_tab):
            window._set_current_main_page(tab)
            app.processEvents()
            if not tab.table.isVisible() or tab.table.rowCount() != 2:
                raise AssertionError("legacy hidden-list preference hid or changed a table")
            buttons = {button.text() for button in tab.findChildren(QPushButton)}
            if {"Скрыть список", "Показать список"} & buttons:
                raise AssertionError("retired list-hiding action remains in the UI")
            tab.search.setText("Smoke Game A")
            QTest.keyClick(tab.search, Qt.Key_Return)
            app.processEvents()
            if tab.selected_game_id() != first_id:
                raise AssertionError("Enter navigation failed after list-hiding removal")
            if window.games_tab.selected_game_id() != first_id:
                raise AssertionError("Enter selection stopped synchronizing with the main list")
            if window.public_tab.selected_game_id() != first_id:
                raise AssertionError("Enter selection stopped synchronizing with Public")
            tab.search.clear()
            app.processEvents()
            if window.size() != original_size:
                raise AssertionError("list tab/search changed the outer window size")

        # The narrow shell reflows tabs while preserving top-level geometry.
        window._set_current_main_page(window.auction_tab)
        narrow_width = window.minimumWidth()
        narrow_height = max(window.minimumHeight() + 80, 440)
        window.resize(narrow_width, narrow_height)
        app.processEvents()
        narrow_size = window.size()
        if narrow_size.width() >= 1100 or narrow_size.height() >= 700:
            raise AssertionError("P06 could not resize below 1100x700")
        auction_scroll = window._main_tab_scroll_by_page[window.auction_tab]
        window._update_responsive_layout()
        app.processEvents()
        if auction_scroll.horizontalScrollBar().maximum() != 0:
            raise AssertionError("P06 narrow Auction page still scrolls horizontally")
        if auction_scroll.horizontalScrollBarPolicy() != Qt.ScrollBarAlwaysOff:
            raise AssertionError("P06 Auction page-level horizontal bar is enabled")

        # The Conduct page collapses the combined operator panes vertically;
        # the table keeps its own columns while the nested page never pans.
        auction = window.auction_tab
        auction.auction_tabs.setCurrentWidget(auction.conduct_page)
        app.processEvents()
        window._update_responsive_layout()
        app.processEvents()
        if auction.conduct_content_layout.direction() != QBoxLayout.TopToBottom:
            raise AssertionError("P06 narrow Auction Conduct panes did not stack")
        if auction.conduct_content_scroll.horizontalScrollBar().maximum() != 0:
            raise AssertionError("P06 Auction Conduct content still scrolls horizontally")
        auction.auction_tabs.setCurrentWidget(auction.lots_page)
        app.processEvents()

        # Compact tabs retain readable short captions and distinct accent colors.
        expected_compact_titles = [
            "Спис", "Публ", "Муз", "Аук", "Ист", "Жур", "OBS", "Наст",
        ]
        if not window._main_tabs_compact:
            raise AssertionError("P06 minimum-width tab strip did not compact")
        if [window.tabs.tabText(index) for index in range(window.tabs.count())] != expected_compact_titles:
            raise AssertionError("P06 compact tab captions are unclear or out of order")
        if any(window.tabs.tabIcon(index).isNull() for index in range(window.tabs.count())):
            raise AssertionError("P06 compact tab strip is missing a colored icon")
        if any(
            not window.tabs.tabToolTip(index)
            for index in range(window.tabs.count())
        ):
            raise AssertionError("P06 compact tabs lost their full-name tooltips")
        compact_icon_colors = []
        for index in range(window.tabs.count()):
            icon_image = window.tabs.tabIcon(index).pixmap(18, 18).toImage()
            pixels = []
            for y in range(icon_image.height()):
                for x in range(icon_image.width()):
                    color = icon_image.pixelColor(x, y)
                    if color.alpha() > 32:
                        rgb = (color.red(), color.green(), color.blue())
                        if max(rgb) - min(rgb) > 20:
                            pixels.append(rgb)
            if not pixels:
                raise AssertionError("P06 compact tab icon rendered white or grayscale")
            compact_icon_colors.append(
                tuple(round(sum(pixel[channel] for pixel in pixels) / len(pixels)) for channel in range(3))
            )
        if len(set(compact_icon_colors)) != window.tabs.count():
            raise AssertionError("P06 compact tab icons do not use distinct colors")
        compact_bar = window.tabs.tabBar()
        if compact_bar.tabRect(window.tabs.count() - 1).right() >= compact_bar.width():
            raise AssertionError("P06 compact tab labels do not fit in the tab strip")

        # Each main workspace must fit the tab viewport; only table widgets
        # may own horizontal column scrolling.
        for page in main_tab_pages:
            window._set_current_main_page(page)
            app.processEvents()
            QTest.qWait(20)
            if page is window.music_tab:
                # First visit must settle its responsive layout without a
                # manual resize or direct call to _update_responsive_layout().
                QTest.qWait(80)
                app.processEvents()
                transport_buttons = (
                    window.music_tab.previous_btn,
                    window.music_tab.stop_btn,
                    window.music_tab.play_pause_btn,
                    window.music_tab.next_btn,
                )
                transport_centers = [
                    button.geometry().center().y()
                    for button in transport_buttons
                ]
                if max(transport_centers) - min(transport_centers) > 4:
                    raise AssertionError(
                        "P06 Music transport buttons stayed vertically stretched "
                        "on first visit"
                    )
                if any(
                    button.width() > button.sizeHint().width() + 20
                    for button in transport_buttons
                ):
                    raise AssertionError(
                        "P06 Music transport buttons filled the row on first visit"
                    )
            else:
                window._update_responsive_layout()
                app.processEvents()
            page_scroll = window._main_tab_scroll_by_page[page]
            if page_scroll.horizontalScrollBar().maximum() != 0:
                raise AssertionError(
                    f"P06 {page_scroll.accessibleName()} page still scrolls horizontally"
                )
            if page.width() > page_scroll.viewport().width() + 1:
                raise AssertionError(
                    f"P06 {page_scroll.accessibleName()} content exceeds its viewport"
                )
        # Music volume controls stay together because they fit the compact row.
        window._set_current_main_page(window.music_tab)
        app.processEvents()
        window._update_responsive_layout()
        app.processEvents()
        music_step_layout = window.music_tab.volume_control.layout()
        if music_step_layout is None or music_step_layout.direction() != QBoxLayout.LeftToRight:
            raise AssertionError("P06 Music volume controls split vertically at minimum width")
        if window.music_tab.volume_control.width() < music_step_layout.sizeHint().width():
            raise AssertionError("P06 Music volume control is clipped at minimum width")

        # Every OBS eyedropper is square, compact, and grouped beside its color field.
        window._set_current_main_page(window.stream_tab)
        app.processEvents()
        window._update_responsive_layout()
        app.processEvents()
        stream_pipettes = [
            button
            for button in window.stream_tab.findChildren(QPushButton)
            if button.property("screenColorPicker")
        ]
        if not stream_pipettes:
            raise AssertionError("P06 OBS screen-color pipettes were not marked for layout QA")
        if any("padding: 0px" not in button.styleSheet() for button in stream_pipettes):
            raise AssertionError("P06 OBS eyedropper glyph is clipped by button padding")
        for pipette in stream_pipettes:
            if (
                pipette.minimumWidth() != pipette.maximumWidth()
                or pipette.minimumHeight() != pipette.maximumHeight()
                or pipette.minimumWidth() != pipette.minimumHeight()
            ):
                raise AssertionError("P06 OBS eyedropper is not a compact square button")
            group = pipette.parentWidget()
            if group is None:
                raise AssertionError("P06 OBS eyedropper has no color-field container")
            contents = group.contentsRect()
            pipette_rect = pipette.geometry()
            if (
                contents.height() < pipette.minimumHeight()
                or pipette_rect.top() < contents.top()
                or pipette_rect.bottom() > contents.bottom()
            ):
                raise AssertionError("P06 OBS eyedropper is clipped vertically by its color row")
            if (
                group.layout() is None
                or group.layout().direction() != QBoxLayout.LeftToRight
            ):
                raise AssertionError("P06 OBS eyedropper is separated from its color field")

        typography_groups = window.stream_tab.typography_control_groups
        if len(typography_groups) != len(window.stream_tab.typography_controls):
            raise AssertionError("P06 OBS typography control groups are incomplete")
        for key, (size_group, color_group) in typography_groups.items():
            if (
                size_group.layout() is None
                or size_group.layout().direction() != QBoxLayout.LeftToRight
                or color_group.layout() is None
                or color_group.layout().direction() != QBoxLayout.LeftToRight
            ):
                raise AssertionError(
                    f"P06 OBS typography controls split inside their groups: {key}"
                )
            size_spin = window.stream_tab.typography_controls[key][1]
            if size_spin.width() < 140:
                raise AssertionError(
                    f"P06 OBS font-size field is squeezed: {key}"
                )
            arrows = [
                button
                for button in size_group.findChildren(QPushButton)
                if button.text() in {"▲", "▼"}
            ]
            if {button.text() for button in arrows} != {"▲", "▼"}:
                raise AssertionError(
                    f"P06 OBS font-size arrows are missing from one row: {key}"
                )
            control_y = [
                size_spin.geometry().center().y(),
                *(button.geometry().center().y() for button in arrows),
            ]
            if max(control_y) - min(control_y) > 4:
                raise AssertionError(
                    f"P06 OBS font-size arrows are not beside the size field: {key}"
                )
            color_button = window.stream_tab.typography_controls[key][2]
            if color_button.width() < 220:
                raise AssertionError(
                    f"P06 OBS typography color field is squeezed: {key}"
                )
            size_labels = [
                label for label in size_group.findChildren(QLabel)
                if label.text() == "Размер:"
            ]
            color_labels = [
                label for label in color_group.findChildren(QLabel)
                if label.text() == "Цвет:"
            ]
            if len(size_labels) != 1 or len(color_labels) != 1:
                raise AssertionError(
                    f"P06 OBS typography group labels are missing: {key}"
                )
            size_label, color_label = size_labels[0], color_labels[0]
            if (
                size_spin.geometry().left()
                - (size_label.geometry().left() + size_label.width())
                > size_group.layout().spacing() + 2
                or color_button.geometry().left()
                - (color_label.geometry().left() + color_label.width())
                > color_group.layout().spacing() + 2
            ):
                raise AssertionError(
                    f"P06 OBS typography labels leave gaps before their fields: {key}"
                )
            for group in (size_group, color_group):
                child_widths = []
                for index in range(group.layout().count()):
                    child = group.layout().itemAt(index).widget()
                    if child is not None:
                        child_widths.append(
                            max(
                                child.minimumWidth(),
                                child.minimumSizeHint().width(),
                                child.sizeHint().width(),
                            )
                        )
                required_width = sum(child_widths) + max(
                    0, group.layout().spacing()
                ) * max(0, len(child_widths) - 1)
                if group.width() < required_width:
                    raise AssertionError(
                        f"P06 OBS typography group is clipped: {key}"
                    )

        # The narrow Settings → Auction image selector and import action remain
        # in one compact row instead of leaving the label isolated.
        window._set_current_main_page(window.settings_tab)
        settings = window.settings_tab
        settings.settings_tabs.setCurrentWidget(settings.auction_page)
        app.processEvents()
        window._update_responsive_layout()
        app.processEvents()
        image_center_y = (
            settings.wheel_center_image_label.mapTo(settings.auction_page, QPoint(0, 0)).y()
            + settings.wheel_center_image_label.height() / 2
        )
        combo_center_y = (
            settings.wheel_center_image_combo.mapTo(settings.auction_page, QPoint(0, 0)).y()
            + settings.wheel_center_image_combo.height() / 2
        )
        file_center_y = (
            settings.add_wheel_center_file_btn.mapTo(settings.auction_page, QPoint(0, 0)).y()
            + settings.add_wheel_center_file_btn.height() / 2
        )
        if (
            settings.wheel_center_select_row.direction() != QBoxLayout.LeftToRight
            or max(image_center_y, combo_center_y, file_center_y)
            - min(image_center_y, combo_center_y, file_center_y) > 8
        ):
            raise AssertionError("P06 Settings → Auction image field did not stay compact and inline")
        if settings.wheel_center_image_combo.width() > 260:
            raise AssertionError("P06 Settings → Auction image selector grew excessively wide")

        # The main-list table keeps its own vertical bar inside the visible
        # workspace while its horizontal bar stays attached to the table.
        window._set_current_main_page(window.games_tab)
        app.processEvents()
        QTest.qWait(100)
        app.processEvents()
        games_scroll = window._main_tab_scroll_by_page[window.games_tab]
        games_scroll.horizontalScrollBar().setValue(0)
        window._update_responsive_layout()
        app.processEvents()
        if games_scroll.horizontalScrollBar().maximum() != 0:
            raise AssertionError("P06 Games page still scrolls as one wide panel")
        if not window.games_tab._compact_controls_enabled:
            raise AssertionError("P06 minimum-width List controls did not reflow")
        if window.games_tab._sorting_actions_layout.direction() != QBoxLayout.TopToBottom:
            raise AssertionError("P06 List sort and preview actions are not stacked")
        toolbar_controls = (
            window.games_tab.sorting_rules_btn,
            window.games_tab.copy_list_overlay_url_btn,
            window.games_tab.open_list_overlay_preview_btn,
            window.games_tab.total_points_label,
            window.games_tab.clear_all_btn,
        )
        toolbar_positions = [
            widget.mapTo(window.games_tab, QPoint(0, 0))
            for widget in toolbar_controls
        ]
        if any(
            toolbar_positions[index + 1].y() <= toolbar_positions[index].y()
            for index in range(len(toolbar_positions) - 1)
        ):
            raise AssertionError("P06 compact List toolbar order changed")
        toolbar_right_edges = [
            widget.mapTo(window.games_tab, QPoint(widget.width(), 0)).x()
            for widget in toolbar_controls
        ]
        if max(toolbar_right_edges) - min(toolbar_right_edges) > 8:
            raise AssertionError("P06 compact List toolbar is not aligned on the right")
        toolbar_buttons = (
            window.games_tab.sorting_rules_btn,
            window.games_tab.copy_list_overlay_url_btn,
            window.games_tab.open_list_overlay_preview_btn,
            window.games_tab.clear_all_btn,
        )
        for previous, following in zip(toolbar_buttons, toolbar_buttons[1:]):
            gap = (
                following.mapTo(window.games_tab, QPoint(0, 0)).y()
                - previous.mapTo(window.games_tab, QPoint(0, 0)).y()
                - previous.height()
            )
            if gap < 6:
                raise AssertionError(
                    f"P06 compact List toolbar buttons have no breathing room: {gap}px"
                )
        left_actions = [
            widget for widget in window.games_tab._action_buttons if widget.isVisible()
        ]
        for previous, following in zip(left_actions, left_actions[1:]):
            gap = (
                following.mapTo(window.games_tab, QPoint(0, 0)).y()
                - previous.mapTo(window.games_tab, QPoint(0, 0)).y()
                - previous.height()
            )
            if gap < 6:
                raise AssertionError(
                    f"P06 compact List action buttons have no breathing room: {gap}px"
                )
        if left_actions:
            left_action_right = max(
                widget.mapTo(window.games_tab, QPoint(widget.width(), 0)).x()
                for widget in left_actions
            )
            toolbar_left = min(point.x() for point in toolbar_positions)
            if toolbar_left <= left_action_right:
                raise AssertionError("P06 List action toolbar did not move to the right column")
        table = window.games_tab.table
        games_viewport = games_scroll.viewport()
        if table.width() > games_viewport.width():
            raise AssertionError("P06 main-list table exceeds its visible viewport width")
        if table.width() < games_viewport.width() - 40:
            raise AssertionError("P06 main-list table does not fill its visible viewport")
        table_bar = table.verticalScrollBar()
        table_bar_right = table.mapTo(
            games_viewport,
            table_bar.geometry().topRight(),
        ).x()
        if table_bar_right >= games_viewport.width():
            raise AssertionError("P06 main-list vertical bar moved outside the visible workspace")

        # Settings keeps its vertical bar at the visible right edge without
        # introducing a workspace-level horizontal scrollbar.
        settings_scroll = window._main_tab_scroll_by_page[window.settings_tab]
        window._set_current_main_page(window.settings_tab)
        app.processEvents()
        settings_scroll.horizontalScrollBar().setValue(0)
        app.processEvents()
        if settings_scroll.horizontalScrollBar().maximum() != 0:
            raise AssertionError("P06 Settings page still scrolls horizontally")
        if settings_scroll.verticalScrollBar().maximum() <= 0:
            raise AssertionError("P06 Settings workspace has no vertical scrolling")
        if not settings_scroll.verticalScrollBar().isVisible():
            raise AssertionError("P06 Settings vertical scrollbar is not visible at the viewport edge")
        window._set_current_main_page(window.public_tab)
        app.processEvents()
        if window.size() != narrow_size:
            raise AssertionError("switching tabs resized the narrow main window")
        narrow_geometry = window.saveGeometry()
        window.resize(original_size)
        window._ui_settings.setValue("main_window/geometry", narrow_geometry)
        window._ui_settings.sync()
        window._restore_ui_state()
        app.processEvents()
        if window.width() >= 1100 or window.height() >= 700:
            raise AssertionError("P06 did not restore saved geometry below 1100x700")
        window.resize(original_size)
        app.processEvents()
        QTest.qWait(100)
        app.processEvents()
        if window._main_tabs_compact:
            raise AssertionError(
                "P06 wide window did not restore tab captions "
                f"(window={window.width()}x{window.height()}, "
                f"tabs={window.tabs.width()}px)"
            )
        if [window.tabs.tabText(index) for index in range(window.tabs.count())] != expected_tab_titles:
            raise AssertionError("P06 wide tab captions were not restored after compact mode")
        window._set_current_main_page(window.public_tab)
        app.processEvents()

        window._save_ui_state()
        reopened = open_ui_settings(paths.data_dir, migrate_native=False)
        if any(reopened.contains(key) for key in obsolete_list_keys):
            raise AssertionError("obsolete list visibility survived or was re-created on save")
        if not reopened.contains("main_window/geometry"):
            raise AssertionError("retiring list visibility removed saved window geometry")
        if reopened.value("main_window/tab_index", -1, type=int) != 1:
            raise AssertionError("retiring list visibility changed saved tab persistence")
        if reopened.value("main_window/tab_key", "", type=str) != "public":
            raise AssertionError("P06 stable tab key was not persisted")
        window._set_current_main_page(window.games_tab)
        app.processEvents()

        # P04-B / BUG-003: ordinary refresh/filter must not invent a new
        # selection. Explicit Enter/focus actions remain allowed to select.
        games_tab = window.games_tab
        if not games_tab._select_game_row(first_id):
            raise AssertionError("P04-B fixture could not select the first game")
        games_tab.apply_stat_filter(STATUS_NOT_PLAYED)
        app.processEvents()
        if games_tab.table.rowCount() != 1:
            raise AssertionError("P04-B status filter fixture changed unexpectedly")
        if games_tab.selected_game_id() is not None:
            raise AssertionError("filter refresh auto-selected a replacement row")
        if games_tab.edit_btn.isEnabled() or games_tab.delete_btn.isEnabled():
            raise AssertionError("no-selection action buttons stayed enabled")
        if games_tab.archive_btn.isVisible() or games_tab.restore_btn.isVisible():
            raise AssertionError("no-selection archive/restore action stayed visible")

        games_tab.apply_stat_filter("all")
        app.processEvents()
        if games_tab.selected_game_id() is not None:
            raise AssertionError("refresh invented a selection after none was selected")

        games_tab.search.setText("Smoke Game A")
        app.processEvents()
        if games_tab.table.rowCount() != 1:
            raise AssertionError("UI-008 typing stopped filtering the main list live")
        QTest.keyClick(games_tab.search, Qt.Key_Return)
        app.processEvents()
        if games_tab.selected_game_id() != first_id:
            raise AssertionError("explicit Enter search stopped selecting its target")
        if not games_tab.table.hasFocus():
            raise AssertionError("UI-008 Enter did not move focus to the main-list table")
        games_tab.search.clear()
        app.processEvents()
        if games_tab.selected_game_id() != first_id:
            raise AssertionError("refresh failed to preserve a still-visible selection")

        empty_point = games_tab.table.viewport().rect().bottomRight()
        if games_tab.table.indexAt(empty_point).isValid():
            raise AssertionError("P04-B fixture has no empty table area to click")
        QTest.mouseClick(
            games_tab.table.viewport(),
            Qt.MouseButton.LeftButton,
            pos=empty_point,
        )
        app.processEvents()
        if games_tab.selected_game_id() is not None:
            raise AssertionError("empty-area click did not clear the current row")
        if games_tab.edit_btn.isEnabled() or games_tab.delete_btn.isEnabled():
            raise AssertionError("empty-area deselection left row actions enabled")
        if games_tab.archive_btn.isVisible() or games_tab.restore_btn.isVisible():
            raise AssertionError("empty-area deselection left archive actions visible")

        # Leave the remainder of the broad smoke in the same explicit-selection
        # state it had before BUG-003 coverage was added.
        if not games_tab._select_game_row(first_id):
            raise AssertionError("P04-B fixture could not restore explicit selection")

        # P04-C / BUG-004: Public keeps explicit selection actions, but a click
        # on empty table space clears the current selection and ordinary refresh
        # must not invent a replacement selection.
        public_tab = window.public_tab
        window._set_current_main_page(public_tab)
        app.processEvents()
        if not public_tab.select_synced_search_result(first_id):
            raise AssertionError("P04-C fixture could not select the public row")
        if public_tab.selected_game_id() != first_id:
            raise AssertionError("P04-C explicit public selection did not stick")

        public_empty_point = public_tab.table.viewport().rect().bottomRight()
        if public_tab.table.indexAt(public_empty_point).isValid():
            raise AssertionError("P04-C fixture has no empty Public table area to click")
        QTest.mouseClick(
            public_tab.table.viewport(),
            Qt.MouseButton.LeftButton,
            pos=public_empty_point,
        )
        app.processEvents()
        if public_tab.selected_game_id() is not None:
            raise AssertionError("Public empty-area click did not clear selection")

        public_tab.refresh()
        app.processEvents()
        if public_tab.selected_game_id() is not None:
            raise AssertionError("Public refresh invented a selection after deselection")

        public_tab.search.setText("Smoke Game A")
        app.processEvents()
        if public_tab.table.rowCount() != 1:
            raise AssertionError("UI-028 Public live-search stopped showing matches")
        QTest.keyClick(public_tab.search, Qt.Key_Return)
        app.processEvents()
        if public_tab.selected_game_id() != first_id:
            raise AssertionError("Public Enter search stopped selecting its target")
        if not public_tab.table.hasFocus():
            raise AssertionError("UI-028 Enter did not focus Public table")
        if window.games_tab.selected_game_id() != first_id:
            raise AssertionError("UI-028 Enter no longer synchronizes the main list")
        public_tab.search.setText("__ui028_no_such_public_record__")
        app.processEvents()
        if public_tab.table.rowCount() != 0:
            raise AssertionError("UI-028 no-match query did not empty Public table")
        ui028_dialog = {}
        original_ui028_info = QMessageBox.information

        def capture_ui028_info(parent, title, message, *args, **kwargs):
            ui028_dialog["title"] = title
            ui028_dialog["message"] = message
            return QMessageBox.Ok

        QMessageBox.information = capture_ui028_info
        try:
            QTest.keyClick(public_tab.search, Qt.Key_Return)
            app.processEvents()
        finally:
            QMessageBox.information = original_ui028_info
        if ui028_dialog.get("title") != "Поиск" or "ничего не найдено" not in (
            ui028_dialog.get("message") or ""
        ):
            raise AssertionError(f"UI-028 Enter no-result dialog changed: {ui028_dialog!r}")
        public_tab.search.clear()
        app.processEvents()
        if public_tab.table.rowCount() != 2:
            raise AssertionError("UI-028 clearing query failed to restore Public list")

        if not public_tab.select_synced_search_result(first_id):
            raise AssertionError("P04-C fixture could not restore explicit Public selection")
        window._set_current_main_page(window.games_tab)
        app.processEvents()

        # P04-D / UI-007: non-empty main-list search ignores the active
        # statistic filter and includes archived normal records, but the
        # selected statistic filter itself remains unchanged and resumes after
        # the search is cleared.
        search_scope_id = db.add_game(
            Game(
                None,
                "P04 Search Archived",
                "2026-03-01",
                7_500,
                0,
                STATUS_PLAYED,
                "P04 archived search fixture",
            )
        )
        db.archive_game(search_scope_id, True)
        games_tab.apply_stat_filter(STATUS_NOT_PLAYED)
        app.processEvents()
        if games_tab.active_filter != STATUS_NOT_PLAYED:
            raise AssertionError("P04-D fixture failed to activate the statistic filter")
        if games_tab.table.rowCount() != 1:
            raise AssertionError("P04-D pre-search filter fixture changed unexpectedly")

        games_tab.search.setText("P04 Search Archived")
        app.processEvents()
        if games_tab.active_filter != STATUS_NOT_PLAYED:
            raise AssertionError("search changed the selected statistic filter")
        if games_tab.table.rowCount() != 1:
            raise AssertionError("search did not ignore the active statistic filter")
        searched_title = games_tab.table.item(0, 2)
        if searched_title is None or searched_title.text() != "P04 Search Archived":
            raise AssertionError("search did not include the archived matching record")
        if not games_tab.filter_buttons[STATUS_NOT_PLAYED].isChecked():
            raise AssertionError("search visually reset the selected statistic filter")

        # UI-009: exercise the actual Qt clear icon, not an invented reset
        # mechanism. The selected statistic filter must survive this click.
        clear_buttons = [
            button for button in games_tab.search.findChildren(QToolButton)
            if button.isVisible() and button.isEnabled()
        ]
        if len(clear_buttons) != 1:
            raise AssertionError(
                f"UI-009 expected one visible native clear button, got {len(clear_buttons)}"
            )
        QTest.mouseClick(clear_buttons[0], Qt.MouseButton.LeftButton)
        app.processEvents()
        if games_tab.search.text() != "":
            raise AssertionError("UI-009 native clear icon did not clear search text")
        if games_tab.active_filter != STATUS_NOT_PLAYED:
            raise AssertionError("clearing search lost the selected statistic filter")
        if not games_tab.filter_buttons[STATUS_NOT_PLAYED].isChecked():
            raise AssertionError("UI-009 native clear reset the selected filter button")
        if games_tab.table.rowCount() != 1:
            raise AssertionError("clearing search did not restore the active statistic filter")
        restored_title = games_tab.table.item(0, 2)
        if restored_title is None or restored_title.text() != "Smoke Game B":
            raise AssertionError("clearing search restored the wrong filtered result")

        # The Archive filter follows the same established search/clear path.
        games_tab.apply_stat_filter("archive")
        app.processEvents()
        if games_tab.table.rowCount() != 1:
            raise AssertionError("UI-009 archive-filter fixture changed unexpectedly")
        games_tab.search.setText("Smoke Game B")
        app.processEvents()
        if games_tab.table.rowCount() != 1:
            raise AssertionError("UI-009 live search stopped overriding archive filter")
        clear_buttons = [
            button for button in games_tab.search.findChildren(QToolButton)
            if button.isVisible() and button.isEnabled()
        ]
        if len(clear_buttons) != 1:
            raise AssertionError("UI-009 archive search lost its native clear button")
        QTest.mouseClick(clear_buttons[0], Qt.MouseButton.LeftButton)
        app.processEvents()
        if games_tab.search.text() or games_tab.active_filter != "archive":
            raise AssertionError("UI-009 native clear lost archive filter state")
        if not games_tab.filter_buttons["archive"].isChecked():
            raise AssertionError("UI-009 native clear unchecked Archive statistic filter")
        if games_tab.table.rowCount() != 1:
            raise AssertionError("UI-009 native clear did not restore archive entries")
        archived_title = games_tab.table.item(0, 2)
        if archived_title is None or archived_title.text() != "P04 Search Archived":
            raise AssertionError("UI-009 native clear displayed the wrong archive entry")

        db.delete_game(search_scope_id)
        games_tab.apply_stat_filter("all")
        app.processEvents()
        if games_tab.table.rowCount() != 2:
            raise AssertionError("P04-D fixture cleanup changed the normal list")
        if not games_tab._select_game_row(first_id):
            raise AssertionError("P04-D fixture could not restore explicit selection")

        # UI-008: Enter on a live-search miss retains the existing
        # informational dialog, without bringing back the Find button.
        games_tab.search.setText("__ui008_no_matching_record__")
        app.processEvents()
        if games_tab.table.rowCount() != 0:
            raise AssertionError("UI-008 no-match query did not produce an empty table")
        ui008_info = {}
        original_ui008_information = QMessageBox.information

        def capture_ui008_information(parent, title, message, *args, **kwargs):
            ui008_info["title"] = title
            ui008_info["message"] = message
            return QMessageBox.Ok

        QMessageBox.information = capture_ui008_information
        try:
            QTest.keyClick(games_tab.search, Qt.Key_Return)
            app.processEvents()
        finally:
            QMessageBox.information = original_ui008_information
        if ui008_info.get("title") != "Поиск" or "ничего не найдено" not in (
            ui008_info.get("message") or ""
        ):
            raise AssertionError(f"UI-008 no-match Enter message changed: {ui008_info!r}")
        games_tab.search.clear()
        app.processEvents()
        if games_tab.table.rowCount() != 2:
            raise AssertionError("UI-008 clearing a missing query did not restore rows")

        duplicate_error = DuplicateGameError(1, "QA")
        if str(duplicate_error) != "Запись «QA» уже существует в списке.":
            raise AssertionError(
                f"duplicate fallback wording mismatch: {str(duplicate_error)!r}"
            )

        terminology_sources = {
            PROJECT_ROOT / "streaming_manager/views/main_window.py": (
                "Игр в восстановленной копии:",
            ),
            PROJECT_ROOT / "streaming_manager/db/services.py": (
                "Не удалось создать резервную копию перед очисткой игр.",
                "Ни одна игра не была удалена.",
                "Игры не были удалены.",
            ),
        }
        for source_path, forbidden_terms in terminology_sources.items():
            source_text = source_path.read_text(encoding="utf-8")
            for forbidden in forbidden_terms:
                if forbidden in source_text:
                    raise AssertionError(
                        f"legacy game wording remains in {source_path.name}: {forbidden!r}"
                    )

        if window.games_tab.clear_all_btn.text() != "Очистить список":
            raise AssertionError(
                f"clear-list button label mismatch: {window.games_tab.clear_all_btn.text()!r}"
            )

        original_messagebox_exec = QMessageBox.exec
        sorting_rules_capture = {}

        def capture_sorting_rules_exec(box):
            sorting_rules_capture["title"] = box.windowTitle()
            sorting_rules_capture["text"] = box.text()
            sorting_rules_capture["informative"] = box.informativeText()
            return QMessageBox.Ok

        QMessageBox.exec = capture_sorting_rules_exec
        try:
            window.games_tab.show_sorting_rules()
        finally:
            QMessageBox.exec = original_messagebox_exec

        if sorting_rules_capture.get("title") != "Правила сортировки":
            raise AssertionError(
                f"sorting-rules window title mismatch: {sorting_rules_capture!r}"
            )
        if sorting_rules_capture.get("text") != "Автоматическая сортировка списка":
            raise AssertionError(
                f"sorting-rules heading mismatch: {sorting_rules_capture!r}"
            )
        sorting_info = sorting_rules_capture.get("informative", "")
        for expected_text in (
            "Сначала список распределяется по статусу:",
            "• В режиме «Всего» сначала идут все записи вне архива",
            "• Архивные записи располагаются отдельным блоком в самом низу",
        ):
            if expected_text not in sorting_info:
                raise AssertionError(
                    f"sorting-rules wording missing {expected_text!r}: {sorting_info!r}"
                )
        for legacy_text in (
            "Сначала игры распределяются по статусу:",
            "• В режиме «Всего» все активные игры идут первыми",
            "• Архивные игры располагаются отдельным блоком в самом низу",
        ):
            if legacy_text in sorting_info:
                raise AssertionError(
                    f"legacy sorting-rules wording remains {legacy_text!r}"
                )

        original_information = QMessageBox.information
        import_help_capture = {}

        def capture_import_help(parent, title, message, *args, **kwargs):
            import_help_capture["title"] = title
            import_help_capture["message"] = message
            return QMessageBox.Ok

        QMessageBox.information = capture_import_help
        try:
            window.games_tab.show_import_csv_help()
        finally:
            QMessageBox.information = original_information

        if import_help_capture.get("title") != "Правила импорта CSV":
            raise AssertionError(
                f"CSV help title mismatch: {import_help_capture!r}"
            )
        import_help = import_help_capture.get("message", "")
        for expected_text in (
            "Обязательный столбец:\n• НАЗВАНИЕ",
            "• ДАТА",
            "• СТАТУС: ПРОХОДИТСЯ, НЕ ИГРАЛ, ИГРАЛ, ПРОЙДЕНО или ЗАБРОШЕНО",
            "Если баллов нет, рекомендуется указать 0.",
            "2. Формат «Название|Баллы»",
            "Файл без заголовка, одна запись в строке:",
            "Для совместимости также принимаются старые названия столбцов:",
            "• НАЗВАНИЕ ИГРЫ",
            "• ДАТА ВЫХОДА",
            "Если в файле найдена ошибка, импорт не применяется частично.",
        ):
            if expected_text not in import_help:
                raise AssertionError(
                    f"CSV help wording missing {expected_text!r}: {import_help!r}"
                )
        for legacy_text in (
            "Обязателен только столбец «НАЗВАНИЕ ИГРЫ».",
            "Для уже существующей игры",
            "Для новой игры",
            "одна игра в строке",
            "Название игры|Баллы",
        ):
            if legacy_text in import_help:
                raise AssertionError(
                    f"legacy CSV help wording remains {legacy_text!r}"
                )

        clear_dialog = DeleteAllGamesDialog(7)
        if clear_dialog.windowTitle() != "Очистить список":
            raise AssertionError(
                f"clear-list dialog title mismatch: {clear_dialog.windowTitle()!r}"
            )
        clear_labels = {label.text() for label in clear_dialog.findChildren(QLabel)}
        if "Будут удалены все записи: 7" not in clear_labels:
            raise AssertionError(f"clear-list heading mismatch: {sorted(clear_labels)!r}")
        expected_warning = (
            "Операция удалит обычные, архивные и временные записи. "
            "Завершённая история аукционов и Журнал сохранятся. Перед удалением "
            "программа автоматически создаст резервную копию текущей базы."
        )
        if expected_warning not in clear_labels:
            raise AssertionError(f"clear-list warning mismatch: {sorted(clear_labels)!r}")
        if DeleteAllGamesDialog.CONFIRM_TEXT != "УДАЛИТЬ ЗАПИСИ":
            raise AssertionError(
                f"clear-list confirm text mismatch: {DeleteAllGamesDialog.CONFIRM_TEXT!r}"
            )
        if clear_dialog.confirm_edit.placeholderText() != "УДАЛИТЬ ЗАПИСИ":
            raise AssertionError(
                f"clear-list placeholder mismatch: {clear_dialog.confirm_edit.placeholderText()!r}"
            )
        if clear_dialog.delete_button.text() != "Удалить записи":
            raise AssertionError(
                f"clear-list destructive button mismatch: {clear_dialog.delete_button.text()!r}"
            )
        clear_dialog.confirm_edit.setText("УДАЛИТЬ ЗАПИСИ ")
        app.processEvents()
        if clear_dialog.delete_button.isEnabled():
            raise AssertionError("clear-list destructive button enabled for inexact confirmation")
        clear_dialog.confirm_edit.setText("УДАЛИТЬ ЗАПИСИ")
        app.processEvents()
        if not clear_dialog.delete_button.isEnabled():
            raise AssertionError("clear-list destructive button not enabled for exact confirmation")
        clear_dialog.confirm_edit.returnPressed.emit()
        app.processEvents()
        if clear_dialog.result() != QDialog.Accepted:
            raise AssertionError("clear-list Enter confirmation semantics changed")
        clear_dialog.deleteLater()

        window.games_tab.clear_all_btn.setText("Очистка…")
        window.games_tab._clear_all_worker_finished()
        if window.games_tab.clear_all_btn.text() != "Очистить список":
            raise AssertionError(
                "clear-list button reverted to legacy text after worker completion"
            )

        add_dialog = GameDialog()
        if add_dialog.windowTitle() != "Добавить":
            raise AssertionError(
                f"add dialog title mismatch: {add_dialog.windowTitle()!r}"
            )

        if add_dialog.title_edit.placeholderText() != "Введите название":
            raise AssertionError(
                f"add dialog title placeholder mismatch: {add_dialog.title_edit.placeholderText()!r}"
            )
        if add_dialog.review_edit.placeholderText() != "Отзыв — можно оставить пустым":
            raise AssertionError(
                f"add dialog review placeholder mismatch: {add_dialog.review_edit.placeholderText()!r}"
            )
        add_dialog_labels = {label.text() for label in add_dialog.findChildren(QLabel)}
        if "Новая игра" in add_dialog_labels:
            raise AssertionError("legacy add dialog heading remains")
        expected_add_labels = {
            "Название:",
            "Дата (необязательно):",
            "Баллы (необязательно):",
            "Кооператив:",
            "Статус:",
            "Отзыв (необязательно):",
        }
        if not expected_add_labels.issubset(add_dialog_labels):
            raise AssertionError(
                f"add dialog field labels mismatch: {sorted(add_dialog_labels)!r}"
            )
        if (
            "Название игры:" in add_dialog_labels
            or "Дата выхода:" in add_dialog_labels
            or "Дата:" in add_dialog_labels
            or "Баллы:" in add_dialog_labels
            or "Отзыв:" in add_dialog_labels
            or "Кооператив (необязательно):" in add_dialog_labels
            or "Статус (необязательно):" in add_dialog_labels
        ):
            raise AssertionError("legacy/incorrect add dialog field labels remain")
        if add_dialog.amount_edit.value() != 0:
            raise AssertionError("optional points default changed from 0")

        def assert_date_caret_contract(dialog, mode):
            # Directly exercise the formatter with a caret in the middle:
            # auto-inserted separators must not force the caret to the end.
            dialog.date_edit.setText("26012010")
            dialog.date_edit.setCursorPosition(4)
            dialog._format_date_while_typing(dialog.date_edit.text())
            if dialog.date_edit.text() != "26.01.2010":
                raise AssertionError(
                    f"{mode} date auto-format changed: {dialog.date_edit.text()!r}"
                )
            if dialog.date_edit.cursorPosition() != 5:
                raise AssertionError(
                    f"{mode} date caret mapping mismatch after auto-format: "
                    f"{dialog.date_edit.cursorPosition()}"
                )

            # Backspace in the middle previously reproduced BUG-002:
            # the old code always jumped to len(formatted).
            dialog.date_edit.setText("26.01.2010")
            dialog.date_edit.setCursorPosition(4)
            QTest.keyClick(dialog.date_edit, Qt.Key.Key_Backspace)
            app.processEvents()
            if dialog.date_edit.cursorPosition() != 3:
                raise AssertionError(
                    f"{mode} date Backspace moved caret unexpectedly: "
                    f"{dialog.date_edit.cursorPosition()}, {dialog.date_edit.text()!r}"
                )
            if dialog.date_edit.cursorPosition() == len(dialog.date_edit.text()):
                raise AssertionError(f"{mode} date Backspace still jumps caret to end")

            # Delete in the middle follows the same ordinary QLineEdit caret rule.
            dialog.date_edit.setText("26.01.2010")
            dialog.date_edit.setCursorPosition(3)
            QTest.keyClick(dialog.date_edit, Qt.Key.Key_Delete)
            app.processEvents()
            if dialog.date_edit.cursorPosition() != 3:
                raise AssertionError(
                    f"{mode} date Delete moved caret unexpectedly: "
                    f"{dialog.date_edit.cursorPosition()}, {dialog.date_edit.text()!r}"
                )
            if dialog.date_edit.cursorPosition() == len(dialog.date_edit.text()):
                raise AssertionError(f"{mode} date Delete still jumps caret to end")

            # Replacing a selection should keep the ordinary collapsed caret
            # directly after the replacement, just like the title field.
            dialog.date_edit.setText("26.01.2010")
            dialog.date_edit.setSelection(3, 2)
            QTest.keyClicks(dialog.date_edit, "12")
            app.processEvents()
            if dialog.date_edit.text() != "26.12.2010":
                raise AssertionError(
                    f"{mode} date selection replacement changed text unexpectedly: "
                    f"{dialog.date_edit.text()!r}"
                )
            if dialog.date_edit.cursorPosition() != 5:
                raise AssertionError(
                    f"{mode} date selection replacement caret mismatch: "
                    f"{dialog.date_edit.cursorPosition()}"
                )

            # Sequential compact input must retain the established auto-dot behavior.
            dialog.date_edit.clear()
            QTest.keyClicks(dialog.date_edit, "26012010")
            app.processEvents()
            if dialog.date_edit.text() != "26.01.2010":
                raise AssertionError(
                    f"{mode} sequential compact date formatting changed: "
                    f"{dialog.date_edit.text()!r}"
                )
            if dialog.date_edit.cursorPosition() != len(dialog.date_edit.text()):
                raise AssertionError(
                    f"{mode} sequential date caret should remain at the end"
                )

        assert_date_caret_contract(add_dialog, "add")
        add_dialog.deleteLater()

        edit_game = Game(
            1,
            "QA",
            None,
            0,
            0,
            STATUS_NOT_PLAYED,
            "",
            updated_at=None,
        )
        edit_dialog = GameDialog(game=edit_game)
        if edit_dialog.windowTitle() != "Изменить":
            raise AssertionError(
                f"edit dialog title mismatch: {edit_dialog.windowTitle()!r}"
            )
        edit_dialog_labels = {label.text() for label in edit_dialog.findChildren(QLabel)}
        if "Редактирование записи" in edit_dialog_labels:
            raise AssertionError("legacy edit dialog heading remains")
        expected_edit_labels = {
            "Название:",
            "Дата (необязательно):",
            "Баллы (необязательно):",
            "Кооператив:",
            "Статус:",
            "Отзыв (необязательно):",
        }
        if not expected_edit_labels.issubset(edit_dialog_labels):
            raise AssertionError(
                f"edit dialog field labels mismatch: {sorted(edit_dialog_labels)!r}"
            )
        if (
            "Название игры:" in edit_dialog_labels
            or "Дата выхода:" in edit_dialog_labels
            or "Дата:" in edit_dialog_labels
            or "Баллы:" in edit_dialog_labels
            or "Отзыв:" in edit_dialog_labels
            or "Кооператив (необязательно):" in edit_dialog_labels
            or "Статус (необязательно):" in edit_dialog_labels
        ):
            raise AssertionError("legacy/incorrect edit dialog field labels remain")
        assert_date_caret_contract(edit_dialog, "edit")
        edit_dialog.deleteLater()
        app.processEvents()

        # P01 shell contract: no top-level menus remain, the replacement hint
        # is informational only, and exactly one direct MainWindow F5 action
        # still performs the existing eager refresh_all() path.
        menu_actions = window.menuBar().actions()
        if menu_actions:
            raise AssertionError(
                f"unexpected top-level menu actions remain: "
                f"{[action.text() for action in menu_actions]}"
            )
        if window.refresh_hint_label.text() != "F5 — обновить данные во всех разделах":
            raise AssertionError(
                f"refresh hint mismatch: {window.refresh_hint_label.text()!r}"
            )
        if not window.refresh_hint_label.testAttribute(Qt.WA_TransparentForMouseEvents):
            raise AssertionError("refresh hint became interactive")

        f5_actions = [
            action
            for action in window.actions()
            if action.shortcut() == QKeySequence("F5")
        ]
        if len(f5_actions) != 1:
            raise AssertionError(
                f"expected one MainWindow F5 action, found {len(f5_actions)}"
            )

        original_games_refresh = window.games_tab.refresh
        f5_refresh_calls = 0

        def counted_games_refresh():
            nonlocal f5_refresh_calls
            f5_refresh_calls += 1
            return original_games_refresh()

        window.games_tab.refresh = counted_games_refresh
        try:
            f5_actions[0].trigger()
            app.processEvents()
        finally:
            window.games_tab.refresh = original_games_refresh
        if f5_refresh_calls != 1:
            raise AssertionError(
                f"F5 did not call refresh_all() exactly once: "
                f"Games refresh count={f5_refresh_calls}"
            )

        auction_tabs = [
            window.auction_tab.auction_tabs.tabText(i)
            for i in range(window.auction_tab.auction_tabs.count())
        ]
        if auction_tabs != ["Лоты", "Проведение"]:
            raise AssertionError(f"auction tabs mismatch: {auction_tabs}")

        settings_tabs = [
            window.settings_tab.settings_tabs.tabText(i)
            for i in range(window.settings_tab.settings_tabs.count())
        ]
        if settings_tabs != ["Общие", "Аукцион", "Интеграции", "Конвертация"]:
            raise AssertionError(f"settings tabs mismatch: {settings_tabs}")
        if hasattr(window.settings_tab, "export_page"):
            raise AssertionError("removed Settings Export page is still constructed")
        p06_general_page = window.settings_tab.general_page
        p06_general_labels = [
            label.text() for label in p06_general_page.findChildren(QLabel)
        ]
        if "Резервная копия базы данных" not in p06_general_labels:
            raise AssertionError("P06 database-backup section heading is missing")
        db_path_label = next(
            (
                label
                for label in p06_general_page.findChildren(QLabel)
                if label.text().startswith("База данных:\n")
            ),
            None,
        )
        if db_path_label is None:
            raise AssertionError("P06 current database path is not visible")
        if not (db_path_label.textInteractionFlags() & Qt.TextSelectableByMouse):
            raise AssertionError("P06 current database path is not selectable")
        if "Полная резервная копия программы" not in p06_general_labels:
            raise AssertionError("P06 full-backup section is not visually separated")
        if (
            "Полная резервная копия предназначена для восстановления данных после "
            "полного удаления и повторной установки InOneLine."
            not in p06_general_labels
        ):
            raise AssertionError("P06 full-backup explanation is missing")
        if not any(
            label.startswith(
                "Защищённые данные подключений предназначены для восстановления "
            )
            for label in p06_general_labels
        ):
            raise AssertionError("P06 DPAPI limitation is not shown in advance")
        if "Интеграции" in p06_general_labels:
            raise AssertionError("P06 redundant Integrations explainer remains")
        settings_titles = [
            window.settings_tab.settings_tabs.tabText(index)
            for index in range(window.settings_tab.settings_tabs.count())
        ]
        if "Интеграции" not in settings_titles:
            raise AssertionError("P06 removed the real Integrations settings tab")
        p06_backup_buttons = {
            button.text() for button in p06_general_page.findChildren(QPushButton)
        }
        if "Создать полную резервную копию…" not in p06_backup_buttons:
            raise AssertionError("P06 full-backup button caption mismatch")
        if "Восстановить из полной резервной копии…" not in p06_backup_buttons:
            raise AssertionError("P06 full-restore button caption mismatch")
        if not {"Открыть папку данных", "Открыть папку резервных копий"}.issubset(
            p06_backup_buttons
        ):
            raise AssertionError("P06 data-folder button captions mismatch")
        if hasattr(window.settings_tab, "export_public_csv_btn"):
            raise AssertionError("duplicate public export buttons remain in Settings")
        if hasattr(window.auction_tab, "legacy_export_page") or hasattr(
            window.auction_tab, "export_rules_btn"
        ):
            raise AssertionError("legacy compatible export page is still constructed")
        if any(
            button.text() == "Копировать список"
            for button in window.auction_tab.findChildren(QPushButton)
        ):
            raise AssertionError("legacy copy-list action remains visible in Auction")

        expected_public_export_labels = ["Экспорт CSV", "Экспорт JSON", "Экспорт Excel"]
        actual_public_export_labels = [
            button.text()
            for button in window.public_tab.findChildren(QPushButton)
            if button.text() in expected_public_export_labels
        ]
        if sorted(actual_public_export_labels) != sorted(expected_public_export_labels):
            raise AssertionError(
                f"Public export actions are missing or duplicated: {actual_public_export_labels!r}"
            )
        for attr in (
            "public_xlsx_create_btn",
            "public_xlsx_connect_btn",
            "public_xlsx_disconnect_btn",
        ):
            if not isinstance(getattr(window.public_tab, attr, None), QPushButton):
                raise AssertionError(f"P1 Public XLSX mirror control missing: {attr}")
        if getattr(window.games_tab, "shared_xlsx_section", None) is None:
            raise AssertionError("shared main-list XLSX mirror was not re-hosted on Games")
        temporary_xlsx_host = getattr(
            window.auction_tab, "_shared_xlsx_controls_host", None
        )
        if temporary_xlsx_host is None or not temporary_xlsx_host.isHidden():
            raise AssertionError("empty AuctionTab XLSX controls host is not hidden")

        pool = QThreadPool.globalInstance()
        if pool.maxThreadCount() > 8:
            raise AssertionError(f"global worker pool unbounded: {pool.maxThreadCount()}")
        if pool.expiryTimeout() != 10_000:
            raise AssertionError(f"worker expiry timeout changed: {pool.expiryTimeout()}")

        print("[2/4] Table reuse / conditional UI")
        window.games_tab.refresh()
        window.public_tab.refresh()
        app.processEvents()
        games_item = window.games_tab.table.item(0, 2)
        public_item = window.public_tab.table.item(0, 1)
        if games_item is None or public_item is None:
            raise AssertionError("table reuse smoke has no initial cells")
        window.games_tab.refresh()
        window.public_tab.refresh()
        app.processEvents()
        if window.games_tab.table.item(0, 2) is not games_item:
            raise AssertionError("Games refresh replaced an unchanged Qt cell")
        if window.public_tab.table.item(0, 1) is not public_item:
            raise AssertionError("Public refresh replaced an unchanged Qt cell")

        games = window.games_tab
        games._apply_action_state(SimpleNamespace(archived=False), 1)
        if games.archive_btn.isHidden() or not games.restore_btn.isHidden():
            raise AssertionError("1.0.8 active-game conditional actions are wrong")
        games._apply_action_state(SimpleNamespace(archived=True), 1)
        if not games.archive_btn.isHidden() or games.restore_btn.isHidden():
            raise AssertionError("1.0.8 archived-game conditional actions are wrong")
        games._apply_action_state(None, 0)
        if not games.archive_btn.isHidden() or not games.restore_btn.isHidden():
            raise AssertionError("1.0.8 no-selection actions are visible")

        stream = window.stream_tab
        stream.timer_background_transparent.setChecked(True)
        stream._update_timer_background_enabled_state()
        if not stream.timer_background_color_row.isHidden():
            raise AssertionError("Timer color row visible in transparent mode")
        stream.timer_background_color_mode.setChecked(True)
        stream._update_timer_background_enabled_state()
        if stream.timer_background_color_row.isHidden():
            raise AssertionError("Timer color row hidden in color mode")

        auto = window.settings_tab
        auto.auto_extend_leader_enabled.setChecked(False)
        auto._update_auto_extend_conditional_visibility()
        if not auto.auto_extend_leader_duration.isHidden():
            raise AssertionError("auto-extend duration visible while reason is disabled")
        auto.auto_extend_leader_enabled.setChecked(True)
        auto._update_auto_extend_conditional_visibility()
        if auto.auto_extend_leader_duration.isHidden():
            raise AssertionError("auto-extend duration hidden while reason is enabled")

        public = window.public_tab
        public._public_xlsx_path = None
        public._public_xlsx_update_controls()
        if not public.public_xlsx_disconnect_btn.isHidden():
            raise AssertionError("Public XLSX disconnect visible while disconnected")

        print("[3/4] D26 soundtrack selectors / live gain")
        auction = window.auction_tab
        position_fixture = [{
            "game_id": 9_999_001,
            "title": "P07 Table Fixture",
            "start_position": 2,
            "current_position": 3,
            "total_sm_points": 1200,
            "auction_only": False,
            "review": "",
        }]
        auction._populate_lot_table(
            auction.table,
            position_fixture,
            None,
            position_mode=POSITION_COLUMNS_SINGLE,
        )
        if auction.table.columnCount() != 3 or [
            auction.table.horizontalHeaderItem(column).text()
            for column in range(auction.table.columnCount())
        ] != ["ПОЗИЦИЯ", "НАЗВАНИЕ", "БАЛЛЫ"]:
            raise AssertionError("Auction single-position table schema is wrong")
        if auction.table.item(0, 0).text() != "3":
            raise AssertionError("Auction current position value is wrong")
        if auction._lot_compact_columns != (0, 2):
            raise AssertionError("Auction single-position compact columns are wrong")

        auction._populate_lot_table(
            auction.conduct_table,
            position_fixture,
            None,
            position_mode=POSITION_COLUMNS_SINGLE,
        )
        auction._sync_wheel_chance_visibility(None)
        if (
            auction.conduct_table.columnCount() != 4
            or getattr(auction, "_conduct_chance_column", None) != 2
            or not auction.conduct_table.isColumnHidden(2)
        ):
            raise AssertionError("Auction compact chance-column layout is wrong")

        auction._populate_lot_table(
            auction.table,
            position_fixture,
            None,
            position_mode=POSITION_COLUMNS_START_CURRENT,
        )
        if (
            auction.table.columnCount() != 4
            or [
                auction.table.horizontalHeaderItem(column).text()
                for column in range(auction.table.columnCount())
            ]
            != ["СТАРТ", "ТЕКУЩАЯ", "НАЗВАНИЕ", "БАЛЛЫ"]
        ):
            raise AssertionError("Auction start/current table schema is wrong")
        if [
            auction.table.item(0, column).text()
            for column in (0, 1)
        ] != ["2", "3"]:
            raise AssertionError("Auction start/current position values are wrong")
        if auction._lot_compact_columns != (0, 1, 3):
            raise AssertionError("Auction start/current compact columns are wrong")

        auction._populate_lot_table(
            auction.conduct_table,
            position_fixture,
            None,
            position_mode=POSITION_COLUMNS_START_CURRENT,
        )
        auction._sync_wheel_chance_visibility(None)
        if (
            auction.conduct_table.columnCount() != 5
            or getattr(auction, "_conduct_chance_column", None) != 3
            or not auction.conduct_table.isColumnHidden(3)
        ):
            raise AssertionError("Auction start/current chance-column layout is wrong")
        if auction._conduct_compact_columns != (0, 1, 3, 4):
            raise AssertionError("Auction start/current conduct compact columns are wrong")
        window._finalize_compact_headers_after_polish()
        auction.refresh()

        if int(auction.wheel_soundtrack_combo.currentData() or 0) != wheel_asset.id:
            raise AssertionError("wheel soundtrack selection was not restored")
        if auction.wheel_soundtrack_volume.value() != 37:
            raise AssertionError("wheel soundtrack volume was not restored")
        if "доступен" not in auction.wheel_soundtrack_status.text().lower():
            raise AssertionError("healthy wheel soundtrack is not available")

        if int(auction.auction_soundtrack_combo.currentData() or 0) != auction_asset.id:
            raise AssertionError("auction soundtrack selection was not restored")
        if auction.auction_soundtrack_volume.value() != 29:
            raise AssertionError("auction soundtrack volume was not restored")
        if "доступен" not in auction.auction_soundtrack_status.text().lower():
            raise AssertionError("healthy auction soundtrack is not available")

        auction.wheel_soundtrack_volume.setValue(41)
        auction.wheel_soundtrack_mute.setChecked(True)
        app.processEvents()
        if db.get_setting(WHEEL_SOUNDTRACK_VOLUME_KEY) != "41":
            raise AssertionError("wheel volume did not persist")
        if db.get_setting(WHEEL_SOUNDTRACK_MUTE_KEY) != "1":
            raise AssertionError("wheel mute did not persist")
        if auction.wheel_audio.volume_percent != 41 or not auction.wheel_audio.muted:
            raise AssertionError("wheel live gain/mute did not reach audio engine")
        auction.wheel_soundtrack_mute.setChecked(False)

        auction.auction_soundtrack_volume.setValue(33)
        auction.auction_soundtrack_mute.setChecked(True)
        app.processEvents()
        if db.get_setting(AUCTION_SOUNDTRACK_VOLUME_KEY) != "33":
            raise AssertionError("auction volume did not persist")
        if db.get_setting(AUCTION_SOUNDTRACK_MUTE_KEY) != "1":
            raise AssertionError("auction mute did not persist")
        if auction.auction_audio.volume_percent != 33 or not auction.auction_audio.muted:
            raise AssertionError("auction live gain/mute did not reach audio engine")
        auction.auction_soundtrack_mute.setChecked(False)

        moved_track = root / "wheel-smoke-moved.wav"
        wheel_track.replace(moved_track)
        auction.refresh_force()
        app.processEvents()
        if auction.wheel_soundtrack_combo.currentData():
            raise AssertionError("missing external wheel soundtrack stayed selected")
        if db.get_setting(WHEEL_SOUNDTRACK_MEDIA_ID_KEY, ""):
            raise AssertionError("missing external wheel soundtrack setting was not cleared")
        if "не выбран" not in auction.wheel_soundtrack_status.text().lower():
            raise AssertionError("missing external wheel soundtrack did not fall back safely")

        replacement = db.register_external_media_asset(MEDIA_CATEGORY_SOUNDTRACK, moved_track)
        auction._adopt_wheel_soundtrack_asset(replacement.id)
        app.processEvents()
        if int(auction.wheel_soundtrack_combo.currentData() or 0) != replacement.id:
            raise AssertionError("re-selected external wheel soundtrack was not adopted")
        if "доступен" not in auction.wheel_soundtrack_status.text().lower():
            raise AssertionError("re-selected wheel soundtrack stayed unavailable")

        assert_p08_background_media(app, window, root)
        assert_shared_projection_preserves_external_edit(app, window, root)

        def check_lifecycle_tables(phase, expected):
            auction.refresh_force()
            window._finalize_compact_headers_after_polish()
            expected_headers = (
                ["СТАРТ", "ТЕКУЩАЯ"]
                if expected == POSITION_COLUMNS_START_CURRENT
                else ["ПОЗИЦИЯ", "НАЗВАНИЕ"]
            )
            for table in (auction.table, auction.conduct_table):
                actual_headers = [
                    table.horizontalHeaderItem(column).text()
                    for column in range(2)
                ]
                if actual_headers != expected_headers:
                    raise AssertionError(
                        f"P07 {phase}: local table headers "
                        f"{actual_headers!r} != {expected_headers!r}"
                    )

        assert_elimination_lifecycle(db, check_lifecycle_tables)
        print("P07_LOCAL_AND_OVERLAY_ELIMINATION_LIFECYCLE=PASS")

        # Populate the temporary database so the main-list row scrollbar is
        # genuinely visible, then verify it stays inside the viewport at x=0.
        window.games_tab.search.clear()
        window.games_tab.active_filter = "all"
        for index in range(80):
            db.add_game(
                Game(
                    None,
                    f"Scrollbar Fixture {index:02d}",
                    None,
                    0,
                    0,
                    STATUS_NOT_PLAYED,
                    "",
                )
            )
        window.games_tab.refresh()
        window._set_current_main_page(window.games_tab)
        window.resize(
            window.minimumWidth(),
            max(window.minimumHeight() + 80, 440),
        )
        app.processEvents()
        QTest.qWait(100)
        app.processEvents()
        games_scroll = window._main_tab_scroll_by_page[window.games_tab]
        games_scroll.horizontalScrollBar().setValue(0)
        window._update_responsive_layout()
        app.processEvents()
        main_table = window.games_tab.table
        main_vertical_bar = main_table.verticalScrollBar()
        if games_scroll.horizontalScrollBar().maximum() != 0:
            raise AssertionError("P06 narrow main-list page still has horizontal scrolling")
        if main_table.horizontalScrollBar().maximum() <= 0:
            raise AssertionError("P06 narrow main-list table has no column scrolling")
        main_table.horizontalScrollBar().setValue(
            main_table.horizontalScrollBar().maximum()
        )
        app.processEvents()
        if main_vertical_bar.maximum() <= 0 or not main_vertical_bar.isVisible():
            raise AssertionError("P06 narrow main-list vertical scrollbar is not usable")
        main_bar_right = main_table.mapTo(
            games_scroll.viewport(),
            main_vertical_bar.geometry().topRight(),
        ).x()
        if main_bar_right >= games_scroll.viewport().width():
            raise AssertionError("P06 vertical row scrollbar moved off-screen when columns scroll")

        print("[4/4] Clean shutdown")
        window.close()
        app.processEvents()
        window = None
        QTest.qWait(450)
        QThreadPool.globalInstance().waitForDone(30_000)
        app.processEvents()
        gc.collect()

        if uncaught:
            exc_type, exc_value = uncaught[0]
            raise AssertionError(
                f"unhandled GUI exception: {exc_type.__name__}: {exc_value}"
            )
    finally:
        sys.excepthook = original_excepthook
        if window is not None:
            try:
                window.close()
                app.processEvents()
            except Exception:
                pass
        try:
            QTest.qWait(450)
            QThreadPool.globalInstance().waitForDone(30_000)
            app.processEvents()
        except Exception:
            pass
        cleanup_tree(root)

    print("1.0.8 GUI REGRESSION CORE: OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
