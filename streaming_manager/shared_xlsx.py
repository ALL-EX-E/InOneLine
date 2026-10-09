from __future__ import annotations

import hashlib
import io
import json
import os
import zipfile
from datetime import date, datetime
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .constants import COOP_FROM_LABEL, COOP_LABELS, STATUS_FROM_LABEL, STATUS_LABELS, STATUS_NOT_PLAYED
from .database import Database, display_date, normalize_title_key, parse_date, points_from_text


SHARED_XLSX_HEADERS = (
    "ПОЗИЦИЯ",
    "НАЗВАНИЕ ИГРЫ",
    "ДАТА ВЫХОДА",
    "БАЛЛЫ",
    "КООП/НЕ КООП",
    "СТАТУС",
    "ОТЗЫВ",
    "АРХИВ",
)
SHARED_XLSX_REQUIRED_HEADERS = SHARED_XLSX_HEADERS[1:]
SHARED_XLSX_SHEET = "spisok"


class SharedXlsxError(RuntimeError):
    """A stable user-facing validation/synchronization error."""


class SharedXlsxTransientError(SharedXlsxError):
    """A temporary filesystem/update state that should be retried."""


def _archive_label(value: int | bool) -> str:
    return "ДА" if bool(value) else "НЕТ"


def _parse_archive(value, *, row_no: int) -> int:
    if value is None or str(value).strip() == "":
        return 0
    if isinstance(value, bool):
        return 1 if value else 0
    text = str(value).strip().upper()
    if text in {"ДА", "YES", "TRUE", "1", "АРХИВ"}:
        return 1
    if text in {"НЕТ", "NO", "FALSE", "0"}:
        return 0
    raise SharedXlsxError(
        f"Некорректное значение АРХИВ в строке {row_no}: {value!r}. "
        "Допустимо: ДА или НЕТ."
    )


def _parse_coop(value, *, row_no: int, title: str) -> int:
    if value is None or str(value).strip() == "":
        return 0
    if isinstance(value, bool):
        return 1 if value else 0
    key = str(value).strip().upper()
    if key in COOP_FROM_LABEL:
        return int(COOP_FROM_LABEL[key])
    raise SharedXlsxError(
        f"Неизвестное значение КООП/НЕ КООП у «{title}» "
        f"(строка {row_no}): {value!r}."
    )


def _parse_status(value, *, row_no: int, title: str) -> str:
    if value is None or str(value).strip() == "":
        return STATUS_NOT_PLAYED
    key = str(value).strip().upper()
    if key in STATUS_FROM_LABEL:
        return str(STATUS_FROM_LABEL[key])
    raise SharedXlsxError(
        f"Неизвестный статус у «{title}» (строка {row_no}): {value!r}."
    )


def _parse_release_date(value, *, row_no: int, title: str) -> str | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    try:
        return parse_date(str(value).strip())
    except ValueError as exc:
        raise SharedXlsxError(
            f"Некорректная дата у «{title}» (строка {row_no}): {value!r}."
        ) from exc


def main_games_rows(db: Database) -> list[dict]:
    """Return the permanent main list and its authoritative effective positions."""
    snapshot = db.games_refresh_snapshot(include_archived=True)
    positions = snapshot["positions"]
    result: list[dict] = []
    for game in snapshot["games"]:
        position = positions.get(int(game.id), (None, None))[1]
        result.append(
            {
                "position": int(position) if position is not None else None,
                "title": game.title,
                "release_date": game.release_date,
                "sm_points": int(game.sm_points),
                "coop": int(game.coop),
                "status": game.status,
                "review": game.review,
                "archived": int(game.archived),
            }
        )
    return result


def state_hash(rows: list[dict]) -> str:
    """Hash logical table content, deliberately ignoring physical row order."""
    canonical = []
    for row in rows:
        canonical.append(
            {
                "title": str(row.get("title") or "").strip(),
                "release_date": row.get("release_date") or None,
                "sm_points": int(row.get("sm_points") or 0),
                "coop": int(row.get("coop") or 0),
                "status": str(row.get("status") or STATUS_NOT_PLAYED),
                "review": str(row.get("review") or "").strip(),
                "archived": int(row.get("archived") or 0),
            }
        )
    canonical.sort(key=lambda row: normalize_title_key(row["title"]))
    payload = json.dumps(canonical, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def presentation_hash(rows: list[dict]) -> str:
    """Hash the complete exported projection, including derived position."""
    canonical = []
    for row in rows:
        position = row.get("position")
        canonical.append(
            {
                "position": int(position) if position is not None else None,
                "title": str(row.get("title") or "").strip(),
                "release_date": row.get("release_date") or None,
                "sm_points": int(row.get("sm_points") or 0),
                "coop": int(row.get("coop") or 0),
                "status": str(row.get("status") or STATUS_NOT_PLAYED),
                "review": str(row.get("review") or "").strip(),
                "archived": int(row.get("archived") or 0),
            }
        )
    canonical.sort(key=lambda row: normalize_title_key(row["title"]))
    payload = json.dumps(canonical, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def position_projection_matches(rows: list[dict], source: dict) -> bool:
    """Check that XLSX's derived position column matches local presentation.

    Position is never imported into database state. This comparison only
    decides whether the connected mirror should be rewritten from the
    authoritative local projection.
    """
    if not bool(source.get("has_position_column")):
        return False
    actual_positions = source.get("position_projection")
    if not isinstance(actual_positions, dict):
        return False

    expected_positions = {
        normalize_title_key(str(row.get("title") or "")): row.get("position")
        for row in rows
    }
    if set(actual_positions) != set(expected_positions):
        return False

    for title_key, expected in expected_positions.items():
        raw = actual_positions[title_key]
        if raw is None or str(raw).strip() == "":
            actual = None
        else:
            try:
                numeric = float(raw)
                if not numeric.is_integer():
                    return False
                actual = int(numeric)
            except (TypeError, ValueError, OverflowError):
                return False
        if actual != (int(expected) if expected is not None else None):
            return False
    return True


def write_shared_xlsx(db: Database, path: str | Path) -> dict:
    """Write a normal Google-Sheets-readable XLSX and atomically replace target."""
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)

    rows = main_games_rows(db)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHARED_XLSX_SHEET
    sheet.append(list(SHARED_XLSX_HEADERS))

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in rows:
        release = date.fromisoformat(row["release_date"]) if row.get("release_date") else None
        sheet.append(
            [
                row.get("position"),
                row["title"],
                release,
                int(row["sm_points"]),
                COOP_LABELS[int(row["coop"])],
                STATUS_LABELS[str(row["status"])],
                row["review"],
                _archive_label(row["archived"]),
            ]
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:H{max(1, sheet.max_row)}"
    widths = {
        "A": 12, "B": 42, "C": 16, "D": 14,
        "E": 18, "F": 18, "G": 55, "H": 12,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        row[2].number_format = "DD.MM.YYYY"
        row[3].number_format = "0"
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Simple dropdowns remain ordinary XLSX validation and are readable by
    # Excel/Google Sheets. They do not create a proprietary file format.
    max_validation_row = max(1000, sheet.max_row + 250)
    coop_validation = DataValidation(type="list", formula1='"НЕ КООП,КООП"', allow_blank=True)
    status_values = ",".join(STATUS_LABELS[value] for value in STATUS_LABELS)
    status_validation = DataValidation(type="list", formula1=f'"{status_values}"', allow_blank=True)
    archive_validation = DataValidation(type="list", formula1='"НЕТ,ДА"', allow_blank=True)
    sheet.add_data_validation(coop_validation)
    sheet.add_data_validation(status_validation)
    sheet.add_data_validation(archive_validation)
    coop_validation.add(f"E2:E{max_validation_row}")
    status_validation.add(f"F2:F{max_validation_row}")
    archive_validation.add(f"H2:H{max_validation_row}")

    temporary = path.with_name(f".{path.stem}.inone-{os.getpid()}-{uuid4().hex}.tmp.xlsx")
    try:
        workbook.save(temporary)
        # Force the completed ZIP container to stable storage before the atomic
        # rename. This also keeps readers from observing a half-written workbook.
        with temporary.open("rb+") as stream:
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    except (PermissionError, OSError) as exc:
        raise SharedXlsxTransientError(
            f"Не удалось записать таблицу «{path.name}». Возможно, файл сейчас открыт другой программой."
        ) from exc
    finally:
        if temporary.exists():
            try:
                temporary.unlink()
            except OSError:
                pass

    stat = path.stat()
    return {
        "path": str(path),
        "hash": state_hash(rows),
        "presentation_hash": presentation_hash(rows),
        "rows": len(rows),
        "signature": (int(stat.st_mtime_ns), int(stat.st_size)),
        "mtime": float(stat.st_mtime),
    }


def read_shared_xlsx(path: str | Path) -> dict:
    """Read and fully validate one shared XLSX before any database mutation."""
    path = Path(path)
    try:
        stat_before = path.stat()
        signature_before = (int(stat_before.st_mtime_ns), int(stat_before.st_size))
        raw = path.read_bytes()
    except (FileNotFoundError, PermissionError, OSError) as exc:
        raise SharedXlsxTransientError(f"Таблица временно недоступна: {path}") from exc

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise SharedXlsxTransientError(
            "Файл XLSX ещё не готов к чтению. Будет выполнена повторная попытка."
        ) from exc

    try:
        sheet = workbook[SHARED_XLSX_SHEET] if SHARED_XLSX_SHEET in workbook.sheetnames else workbook.active
        header_row_no = None
        header_map: dict[str, int] = {}
        # Google Sheets exports can omit the worksheet <dimension> element.
        # In openpyxl read-only mode that makes max_row/max_column None, so
        # header discovery must not depend on precomputed worksheet bounds.
        row_iterator = sheet.iter_rows(values_only=True)
        for row_no, cells in enumerate(row_iterator, start=1):
            if row_no > 20:
                break
            normalized = {
                str(value).strip().upper(): index
                for index, value in enumerate(cells)
                if value is not None and str(value).strip()
            }
            if "НАЗВАНИЕ ИГРЫ" in normalized:
                header_row_no = row_no
                header_map = normalized
                break
        if header_row_no is None:
            raise SharedXlsxError("Не найден заголовок НАЗВАНИЕ ИГРЫ.")

        missing = [header for header in SHARED_XLSX_REQUIRED_HEADERS if header not in header_map]
        if missing:
            raise SharedXlsxError(
                "В совместной таблице отсутствуют обязательные столбцы: " + ", ".join(missing)
            )

        has_position_column = "ПОЗИЦИЯ" in header_map
        position_projection: dict[str, object] = {}
        rows: list[dict] = []
        seen: dict[str, str] = {}
        # Continue the same streaming iterator after the header. This works
        # whether or not the XLSX declares sheet dimensions in advance.
        for row_no, cells in enumerate(row_iterator, start=header_row_no + 1):
            def value(header: str):
                index = header_map[header]
                return cells[index] if index < len(cells) else None

            expected_values = [value(header) for header in SHARED_XLSX_REQUIRED_HEADERS]
            if all(item is None or str(item).strip() == "" for item in expected_values):
                continue

            title_value = value("НАЗВАНИЕ ИГРЫ")
            title = str(title_value or "").strip()
            if not title:
                raise SharedXlsxError(f"Строка {row_no} содержит данные без НАЗВАНИЯ ИГРЫ.")
            title_key = normalize_title_key(title)
            if title_key in seen:
                raise SharedXlsxError(
                    f"Дубликат названия в XLSX: «{title}» (строка {row_no}; "
                    f"ранее встречалось как «{seen[title_key]}»)."
                )
            seen[title_key] = title
            if has_position_column:
                position_projection[title_key] = value("ПОЗИЦИЯ")

            points_value = value("БАЛЛЫ")
            try:
                sm_points = points_from_text(points_value if points_value not in (None, "") else 0)
            except (TypeError, ValueError) as exc:
                raise SharedXlsxError(
                    f"Некорректное значение БАЛЛЫ у «{title}» (строка {row_no}): {points_value!r}."
                ) from exc

            rows.append(
                {
                    "title": title,
                    "release_date": _parse_release_date(value("ДАТА ВЫХОДА"), row_no=row_no, title=title),
                    "sm_points": int(sm_points),
                    "coop": _parse_coop(value("КООП/НЕ КООП"), row_no=row_no, title=title),
                    "status": _parse_status(value("СТАТУС"), row_no=row_no, title=title),
                    "review": str(value("ОТЗЫВ") or "").strip(),
                    "archived": _parse_archive(value("АРХИВ"), row_no=row_no),
                }
            )
    finally:
        workbook.close()

    try:
        stat = path.stat()
        signature_after = (int(stat.st_mtime_ns), int(stat.st_size))
    except OSError as exc:
        raise SharedXlsxTransientError("Таблица изменилась во время чтения; повторная попытка.") from exc
    if signature_after != signature_before:
        raise SharedXlsxTransientError("Таблица изменилась во время чтения; повторная попытка.")

    return {
        "path": str(path),
        "rows": rows,
        "hash": state_hash(rows),
        "has_position_column": has_position_column,
        "position_projection": position_projection,
        "signature": (int(stat.st_mtime_ns), int(stat.st_size)),
        "mtime": float(stat.st_mtime),
    }
