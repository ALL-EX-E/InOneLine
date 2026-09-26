from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .database import Database


PUBLIC_XLSX_HEADERS = (
    "НАЗВАНИЕ ИГРЫ",
    "БАЛЛЫ",
    "ОТЗЫВ",
    "СТАТУС",
)
PUBLIC_XLSX_SHEET = "spisok"


class PublicXlsxError(RuntimeError):
    """Base error for the one-way public-list XLSX mirror."""


class PublicXlsxTransientError(PublicXlsxError):
    """Temporary write/access failure that should be retried."""


def public_mirror_rows(db: Database) -> list[dict]:
    """Return exactly the four public columns in canonical public-list order."""
    return [
        {
            "title": str(row["title"]),
            "sm_points": int(row["sm_points"]),
            "review": str(row["review"] or ""),
            "status": str(row["status"]),
        }
        for row in db.public_games()
    ]


def public_state_hash(rows: list[dict]) -> str:
    """Hash exact public content including row order, but no XLSX presentation."""
    canonical = [
        {
            "title": str(row.get("title") or ""),
            "sm_points": int(row.get("sm_points") or 0),
            "review": str(row.get("review") or ""),
            "status": str(row.get("status") or ""),
        }
        for row in rows
    ]
    payload = json.dumps(
        canonical,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def write_public_xlsx(db: Database, path: str | Path) -> dict:
    """Atomically mirror the current public list to an ordinary XLSX.

    This module is intentionally write-only from the application's perspective:
    external XLSX contents are never read or imported into SQLite.
    """
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
    except (PermissionError, OSError) as exc:
        raise PublicXlsxTransientError(
            f"Не удалось открыть папку для публичной таблицы: {path.parent}"
        ) from exc

    rows = public_mirror_rows(db)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = PUBLIC_XLSX_SHEET
    sheet.append(list(PUBLIC_XLSX_HEADERS))

    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row in rows:
        sheet.append(
            [
                row["title"],
                int(row["sm_points"]),
                row["review"],
                row["status"],
            ]
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:D{max(1, sheet.max_row)}"
    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 55
    sheet.column_dimensions["D"].width = 18

    for cells in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        cells[1].number_format = "0"
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    temporary = path.with_name(
        f".{path.stem}.inone-public-{os.getpid()}-{uuid4().hex}.tmp.xlsx"
    )
    try:
        workbook.save(temporary)
        with temporary.open("rb+") as stream:
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    except (PermissionError, OSError) as exc:
        raise PublicXlsxTransientError(
            f"Ожидание доступа к файлу «{path.name}». "
            "Возможно, он сейчас занят другой программой."
        ) from exc
    finally:
        workbook.close()
        if temporary.exists():
            try:
                temporary.unlink()
            except OSError:
                pass

    try:
        stat = path.stat()
    except OSError as exc:
        raise PublicXlsxTransientError(
            f"Публичная таблица временно недоступна: {path}"
        ) from exc

    return {
        "path": str(path),
        "hash": public_state_hash(rows),
        "rows": len(rows),
        "signature": (int(stat.st_mtime_ns), int(stat.st_size)),
        "mtime": float(stat.st_mtime),
    }
