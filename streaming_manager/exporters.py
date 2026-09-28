from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from .database import Database, format_points


PUBLIC_HEADERS = ["НАЗВАНИЕ ИГРЫ", "БАЛЛЫ", "ОТЗЫВ", "СТАТУС"]


def export_public_csv(db: Database, path: str | Path) -> Path:
    path = Path(path)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(PUBLIC_HEADERS)
        for row in db.public_games():
            writer.writerow([row["title"], int(row["sm_points"]), row["review"], row["status"]])
    return path


def export_public_json(db: Database, path: str | Path) -> Path:
    path = Path(path)
    payload = {
        "version": 1,
        "columns": PUBLIC_HEADERS,
        "games": db.public_games(),
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def export_public_xlsx(db: Database, path: str | Path) -> Path:
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "spisok"
    ws.append(PUBLIC_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in db.public_games():
        ws.append([row["title"], int(row["sm_points"]), row["review"], row["status"]])
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 55
    ws.column_dimensions["D"].width = 18
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = '#,##0'
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(path)
    return path


def export_pointauc_csv(db: Database, path: str | Path) -> Path:
    """Legacy pipe-delimited text: Название|целые баллы SM, без заголовка."""
    path = Path(path)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        for game in db.pointauc_games():
            f.write(f"{game.title}|{format_points(game.sm_points)}\n")
    return path


def pointauc_text(db: Database) -> str:
    return "\n".join(
        f"{game.title}|{format_points(game.sm_points)}" for game in db.pointauc_games()
    )
